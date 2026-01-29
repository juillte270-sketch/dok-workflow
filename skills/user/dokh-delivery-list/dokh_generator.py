import os
import sys
import re
import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Import existing parser setup
project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
sys.path.append(project_root)
from skills.order_processing.scripts.parser import OrderParser
import json

# ========== 설정 ==========
FONT_NAME = '맑은 고딕'
FONT_SIZE = 14
LINE_SPACING = 16  # Reduced from 18 to fit more content

# Global for current client tracking in cell
current_client = ""

def set_cell_border(cell):
    """셀에 검정 테두리 설정"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), '000000')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def get_item_color(item_text, client_name, stock_items, market_items, always_blue):
    """품목 색상 결정"""
    # 새싹은 항상 파란색
    for blue_item in always_blue:
        if blue_item in item_text:
            return RGBColor(0, 0, 255)
    
    # Store check helper
    def check_store_match(target, current):
        return target in current

    # 재고소분 (빨간색)
    for location, items in stock_items.items():
        if check_store_match(location, client_name):
            for item in items:
                # v7 logic: item is explicit "Item Qty(Unit)" or similar
                # We match if item name from list is in text
                pass 
                # Actually, our dynamic extraction gives just item names usually
                if item in item_text:
                    return RGBColor(255, 0, 0)
    
    # 시장소분 (파란색)
    for location, items in market_items.items():
        if check_store_match(location, client_name):
            for item in items:
                if item in item_text:
                    return RGBColor(0, 0, 255)
                    
    return RGBColor(0, 0, 0)

def add_content_to_cell(cell, items, stock_items, market_items, always_blue, section_type=None):
    """셀에 내용 추가"""
    # We need to track current client explicitly for this cell sequence
    global current_client
    
    cell.paragraphs[0].clear()
    
    current_client_local = "" 
    
    # Remove trailing empty strings to save space
    while items and items[-1] == "":
        items.pop()
    
    for i, item in enumerate(items):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(LINE_SPACING)
        
        run = p.add_run(item)
        run.font.name = FONT_NAME
        run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)
        run.font.size = Pt(FONT_SIZE)
        
        # Color Logic
        if item.startswith('-'):
            current_client_local = item[1:].strip()
            run.font.bold = True
            
            # Section headers themselves
            if section_type == 'stock':
                run.font.color.rgb = RGBColor(255, 0, 0)
            elif section_type == 'market':
                run.font.color.rgb = RGBColor(0, 0, 255)
            elif '정복' in item: # v7 logic
                run.font.color.rgb = RGBColor(0, 0, 255)
            else:
                run.font.color.rgb = RGBColor(0, 0, 0)
                
        elif item.startswith('('):
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif item == '':
            pass
        else:
            if section_type == 'stock':
                run.font.color.rgb = RGBColor(255, 0, 0)
            elif section_type == 'market':
                run.font.color.rgb = RGBColor(0, 0, 255)
            elif '정복' in current_client_local:
                run.font.color.rgb = RGBColor(0, 0, 255)
            else:
                run.font.color.rgb = get_item_color(item, current_client_local, stock_items, market_items, always_blue)

def create_page(doc, col1_data, col2_data, col3_data, stock_items, market_items, always_blue, col1_type=None, col2_type=None, col3_type=None):
    """3열 테이블 페이지 생성"""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(6.3)
            set_cell_border(cell)
    
    row = table.rows[0]
    add_content_to_cell(row.cells[0], col1_data, stock_items, market_items, always_blue, section_type=col1_type)
    add_content_to_cell(row.cells[1], col2_data, stock_items, market_items, always_blue, section_type=col2_type)
    add_content_to_cell(row.cells[2], col3_data, stock_items, market_items, always_blue, section_type=col3_type)

class DokhGeneratorV7:
    def __init__(self):
        self.parser = OrderParser()
        self.doc = Document()
        
        # Initial doc setup
        for section in self.doc.sections:
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
            section.left_margin = Cm(0.7)
            section.right_margin = Cm(0.7)
            section.top_margin = Cm(0.7)
            section.bottom_margin = Cm(0.7)
            
        # Hardcoded always blue
        self.always_blue = ['새싹', '무순']

    def load_mappings(self, mappings_path):
        with open(mappings_path, 'r', encoding='utf-8') as f:
            return json.load(f)

    def extract_dynamic_rules(self, orders):
        """Extract stock_items and market_items dicts from the parsed orders"""
        stock_items = {} # {store: [items]}
        market_items = {}
        
        # Search patterns
        pattern = re.compile(r'\(([^)]+)\)')
        
        # 1. Stock Section
        for key, items in orders.items():
            if "재고" in key and "창고" in key:
                for item in items:
                    if item['type'] == 'item':
                        raw_line = item['original_line']
                        match = pattern.search(raw_line)
                        if match:
                            store_hint = match.group(1).split()[0] # Take first word e.g. "마곡"
                            if store_hint not in stock_items: stock_items[store_hint] = []
                            # We store just the item name for matching
                            stock_items[store_hint].append(item['item'])
                            
            if "시장" in key and "구매" in key:
                 for item in items:
                    if item['type'] == 'item':
                        raw_line = item['original_line']
                        match = pattern.search(raw_line)
                        if match:
                            store_hint = match.group(1).split()[0]
                            if store_hint not in market_items: market_items[store_hint] = []
                            market_items[store_hint].append(item['item'])
                            
        return stock_items, market_items

    def flatten_order_to_list(self, orders, store_key):
        """Convert a store's order objects into a list of strings matching v7 format"""
        if store_key not in orders: return []
        
        lines = [f"-{store_key}"]
        for item in orders[store_key]:
            if item['type'] == 'memo':
                lines.append(item['text'])
            else:
                lines.append(f"{item['item']} {item['qty']}{item['unit']}")
        lines.append("") # Empty line after store
        return lines

    def generate(self, input_path, output_path, mappings_path):
        with open(input_path, 'r', encoding='utf-8') as f:
            raw_text = f.read()
            
        orders = self.parser.parse_text(raw_text)
        mappings = self.load_mappings(mappings_path)
        
        stock_items, market_items = self.extract_dynamic_rules(orders)
        
        # Page Config
        is_friday = datetime.datetime.now().weekday() == 4
        config_key = "friday" if is_friday else "default"
        page_config = mappings['page_config'][config_key]
        
        sorted_pages = sorted([k for k in page_config.keys() if k.startswith('page')])
        
        # Prepare "Supplier" and "Stock/Market" remaining lists
        # Identify stores used in pages
        used_stores = []
        for p in sorted_pages:
            if isinstance(page_config[p], list):
                used_stores.extend(page_config[p])
        
        # Identify Supplier sections (not in used_stores, not Stock/Market)
        all_keys = list(orders.keys())
        supplier_keys = []
        stock_key = None
        market_key = None
        
        for k in all_keys:
            # Fuzzy match for used stores
            is_used = False
            for us in used_stores:
                if us.replace(" ", "") in k.replace(" ", ""):
                    is_used = True
                    break
            
            if "재고" in k and "창고" in k:
                stock_key = k
            elif "시장" in k:
                market_key = k
            elif not is_used:
                supplier_keys.append(k)
                
        # --- Generate Pages ---
        for i, page_key in enumerate(sorted_pages):
            content = page_config[page_key]
            
            if i > 0:
                self.doc.add_page_break()
                
            col1 = []
            col2 = []
            col3 = []
            
            c1_type = None
            c2_type = None
            c3_type = None

            if isinstance(content, list):
                # Distribute stores dynamically to balance column heights
                
                # 1. Calculate line counts for each store
                store_data = [] # [(store_name, lines_list, line_count), ...]
                total_lines = 0
                
                # Helper to fuzzy find key
                def get_key(name):
                    for k in orders.keys():
                        if name.replace(" ", "") in k.replace(" ", ""):
                            return k
                    return None

                for s in content:
                    k = get_key(s)
                    if k:
                        lines = self.flatten_order_to_list(orders, k)
                    else:
                        lines = [f"-{s}", "데이터 없음", ""]
                    
                    store_data.append({
                        'name': s,
                        'lines': lines,
                        'count': len(lines)
                    })
                    total_lines += len(lines)
                
                # 2. Target lines per column
                target_per_col = total_lines // 3
                
                # 3. Fill columns sequentially
                c1_lines = []
                c2_lines = []
                c3_lines = []
                
                current_col_lines = 0
                current_col_idx = 0 # 0, 1, 2
                
                for data in store_data:
                    # If current column is full enough, move to next
                    # But if it's the last column, just dump everything
                    if current_col_idx < 2:
                        # Logic: If adding this store exceeds target significantly?
                        # Or just simple fill until >= target
                        if current_col_lines >= target_per_col:
                            current_col_idx += 1
                            current_col_lines = 0
                    
                    if current_col_idx == 0:
                        c1_lines.extend(data['lines'])
                    elif current_col_idx == 1:
                        c2_lines.extend(data['lines'])
                    else:
                        c3_lines.extend(data['lines'])
                    
                    current_col_lines += data['count']

                col1 = c1_lines
                col2 = c2_lines
                col3 = c3_lines

            elif content == "SUPPLIER_LIST":
                # Distribute suppliers dynamically
                
                # 1. Calculate line counts
                supplier_data = [] 
                total_lines = 0
                
                for k in supplier_keys:
                    lines = self.flatten_order_to_list(orders, k)
                    supplier_data.append({
                        'name': k,
                        'lines': lines,
                        'count': len(lines)
                    })
                    total_lines += len(lines)
                
                # 2. Target lines
                target_per_col = total_lines // 3
                
                # 3. Fill columns
                c1 = []
                c2 = []
                c3 = []
                current_lines = 0
                idx = 0
                
                for data in supplier_data:
                    if idx < 2 and current_lines >= target_per_col:
                        idx += 1
                        current_lines = 0
                        
                    if idx == 0: c1.extend(data['lines'])
                    elif idx == 1: c2.extend(data['lines'])
                    else: c3.extend(data['lines'])
                    
                    current_lines += data['count']
                
                col1 = c1
                col2 = c2
                col3 = c3

            elif content == "STOCK_MARKET_LIST":
                # Stock (Col 1), Market (Col 2)
                if stock_key:
                    col1.extend(self.flatten_order_to_list(orders, stock_key))
                    c1_type = 'stock'
                if market_key:
                    col2.extend(self.flatten_order_to_list(orders, market_key))
                    c2_type = 'market'
                    
            create_page(self.doc, col1, col2, col3, stock_items, market_items, self.always_blue, 
                        col1_type=c1_type, col2_type=c2_type, col3_type=c3_type)

        self.doc.save(output_path)
        print("완료!")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        # Defaults for testing
        base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        input_file = os.path.join(base_dir, "data", "inputs", "orders_today.txt")
        mappings_file = os.path.join(base_dir, "skills", "order_processing", "resources", "mappings.json")
        output_file = "test_output_v7.docx"
        
        generator = DokhGeneratorV7()
        generator.generate(input_file, output_file, mappings_file)
    else:
        # Arg mode
        input_file = sys.argv[1]
        output_file = sys.argv[2]
        # derive mappings from known location relative to script
        base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        mappings_file = os.path.join(base_dir, "skills", "order_processing", "resources", "mappings.json")
        
        generator = DokhGeneratorV7()
        generator.generate(input_file, output_file, mappings_file)
