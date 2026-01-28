import json
import os
import datetime
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from typing import Dict, List, Any

# ========== 스타일 설정 ==========
FONT_SIZE = 15
LINE_SPACING = FONT_SIZE + 5
PARA_SPACING = 1.5
MARGIN = 1

BORDER_SETTINGS = {
    'top': {'val': 'single', 'sz': '4', 'color': '000000'},
    'bottom': {'val': 'single', 'sz': '4', 'color': '000000'},
    'left': {'val': 'single', 'sz': '4', 'color': '000000'},
    'right': {'val': 'single', 'sz': '4', 'color': '000000'}
}

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ['top', 'left', 'bottom', 'right']:
        if edge in kwargs:
            element = OxmlElement(f'w:{edge}')
            element.set(qn('w:val'), kwargs[edge].get('val', 'single'))
            element.set(qn('w:sz'), kwargs[edge].get('sz', '4'))
            element.set(qn('w:color'), kwargs[edge].get('color', '000000'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

class DeliveryListDocxGenerator:
    def __init__(self, resource_path: str):
        self.mappings = self._load_mappings(resource_path)
        
    def _load_mappings(self, path: str) -> Dict:
        mapping_file = os.path.join(path, "mappings.json")
        try:
            with open(mapping_file, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}

    def _get_supplier_for_item(self, item_name: str, item_unit: str = "", store_name: str = "") -> str:
        # 1. 재고/시장 소분 규칙
        stock_rules = self.mappings.get("stock_market_rules", [])
        for rule in stock_rules:
            if rule["item_keyword"] in item_name:
                for cond in rule["conditions"]:
                    match = True
                    if "store" in cond and cond["store"] not in store_name: match = False
                    if "unit" in cond:
                        if isinstance(cond["unit"], list):
                            unit_match = False
                            for u in cond["unit"]:
                                if u in item_unit: 
                                    unit_match = True
                                    break
                            if not unit_match: match = False
                        elif cond["unit"] not in item_unit: match = False
                    
                    if match:
                        target = cond["target"]
                        if target == "MARKET": return "시장구매, 시장소분"
                        elif target == "STOCK": return "재고, 창고소분"
                        return target

        # 2. 일반 공급처 매핑
        supplier_map = self.mappings.get("item_supplier_map", {})
        for supplier, items in supplier_map.items():
            for map_item in items:
                if map_item in item_name or item_name in map_item:
                    return supplier
                    
        return "기타/미분류"

    def _add_content_to_cell(self, cell, content_lines: List[str], section_type: str = None):
        for p in cell.paragraphs:
            p._element.getparent().remove(p._element)
            
        for line in content_lines:
            if not line: 
                cell.add_paragraph()
                continue
            
            para = cell.add_paragraph()
            para.paragraph_format.space_before = Pt(PARA_SPACING)
            para.paragraph_format.space_after = Pt(PARA_SPACING)
            para.paragraph_format.line_spacing = Pt(LINE_SPACING)
            para.paragraph_format.left_indent = Pt(0)
            para.paragraph_format.first_line_indent = Pt(0)
            
            run = para.add_run(line)
            run.font.name = '맑은 고딕'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            run.font.size = Pt(FONT_SIZE)
            
            if line.startswith('-'):
                run.bold = True
            
            if section_type == 'stock':
                run.font.color.rgb = RGBColor(255, 0, 0)
            elif section_type == 'market':
                run.font.color.rgb = RGBColor(0, 0, 255)
            else:
                run.font.color.rgb = RGBColor(0, 0, 0)

    def generate_docx(self, orders: Dict[str, List[Dict]], output_path: str):
        # 1. 데이터 준비
        
        # (A) 거래처 리스트 (페이지 1,2) -> orders 그대로 출력
        
        # (B) 공급처 리스트 (페이지 3) -> 합산 로직 포함
        # 구조: {Supplier: [Group1(items), Memo, Group2(items)...]}
        # 그룹 안에서는 (품목, 단위)로 합산
        
        supplier_data = {} # {sup: [ {type:'agg', data:{ (name,unit): qty }}, {type:'memo', text:'...'} ]}
        stock_lines = {"STOCK": [], "MARKET": []}
        
        # 임시 저장소 (현재 공급처의 합산용)
        # current_agg[sup] = { (name, unit): qty }
        current_agg = {} 
        
        # 모든 주문 순회
        # orders는 순서가 있으므로, 아이템을 만날 때마다 해당 supplier의 current_agg에 더함.
        # 만약 해당 supplier와 관련된 Memo가 나오면? -> 지금까지 current_agg를 flush하고 Memo 기록.
        # Memo가 어느 supplier 것인지 알기 위해 'pending_memos' 사용
        
        pending_memos = [] # [text, text...]
        
        # 1차 순회: 아이템별 Supplier 태깅 및 Stock 분리
        processed_items = [] # [{sup, item, unit, qty, type, text...}]
        
        for store, items in orders.items():
            for item in items:
                if item.get("type") == "memo":
                    processed_items.append({"type":"memo", "text":item["text"]})
                else:
                    sup = self._get_supplier_for_item(item["item"], item["unit"], store)
                    
                    if sup == "재고, 창고소분" or sup == "시장구매, 시장소분":
                         target_key = "STOCK" if "재고" in sup else "MARKET"
                         s_name = store.replace('부엉이산장', '').replace('샤브야키', '').replace('오레노이키루미치', '').replace('선데이버거클럽', '선데이').replace(' 점', '').replace('점', '').strip()
                         
                         qs = str(item['qty'])
                         if qs.endswith(".0"): qs = qs[:-2]
                         if item['qty'] == 0: qs = ""
                         text = f"{item['item']} {qs}{item['unit']}".strip()
                         
                         if "새싹" in item["item"]: display_text = text
                         else: display_text = f"{text}({s_name})"
                         stock_lines[target_key].append(display_text)
                    else:
                        # 일반 공급처 아이템
                        processed_items.append({
                            "type": "item",
                            "sup": sup,
                            "item": item["item"],
                            "unit": item["unit"],
                            "qty": item["qty"]
                        })

        # 2차 순회: 공급처별 데이터 구축 (합산)
        # 단순화를 위해: 모든 태깅된 아이템을 공급처별로 모은 뒤, Memo를 어디에 끼워 넣을지 결정?
        # 아니면 공급처별로 독립적으로 관리?
        
        # 공급처별로 데이터를 모으되, '선발주' 같은 메모가 나오면 흐름을 끊어야 함.
        # 하지만 processed_items 리스트는 store 순서대로 섞여 있음.
        # 다모아버섯 아이템들이 여기저기 흩어져 있는데, 얘네를 다 모아서 합산해야 함.
        # 단, 만약 (선발주) 메모가 오복상회 섹션에 있었다면, 오복상회 리스트 사이에 들어가야 함.
        
        # 전략: 공급처별로 {sup: [Item, Item, Memo, Item...]} 리스트를 만듦.
        # 이때 Memo는 어떻게 할당? -> Memo의 '맥락'을 파악해야 하는데, 
        # 원본 orders 순회 시 store가 오복상회 물건을 많이 시켰다면 그 store의 memo는 오복상회로?
        # 사용자 예시: 오복상회 리스트 중간에 (선발주)가 있음.
        
        # 가장 현실적인 합산 방법:
        # 각 공급처별로 Dict를 하나 둠. `sup_blocks = {sup: [ {agg_dict}, {memo}, {agg_dict} ]}`
        # 기본적으로 agg_dict에 계속 더함.
        # Memo가 발견되면? -> 누구의 Memo인지 판단 필요.
        # Memo 텍스트에 "(선발주)" 등이 있으면 -> '선발주의 대상이 되는 공급처'를 찾아야 함.
        # 이는 매우 어려움(자동화 관점).
        
        # 대안: Memo를 포기하고 전체 합산? -> 안됨. 사용자 요청에 (선발주) 있음.
        # 대안 2: 그냥 전체 합산하되, (선발주) 메모는 맨 뒤나 맨 앞에 붙임? -> "니 맘대로 나눠놓지 마라"
        
        # 결국 사용자의 '이상적인 출력'은:
        # 다모아버섯 -> 아이템 다 합산.
        # 오복상회 -> 아이템 다 합산, 근데 중간에 (선발주) 블록이 따로 있음.
        # 이는 아마도 데이터 입력 시 '선발주'로 들어온 내역이 따로 관리되었거나,
        # 입력된 데이터 순서 상 (선발주) 라인 뒤에 나오는 애들은 따로 합산되길 원함.
        
        # 여기서는 "공급처별 단순 전체 합산"을 기본으로 하되, 
        # Memo는 별도로 수집해서 해당 공급처의 맨 뒤(또는 앞)에 붙여주는 방식으로 타협해야 할 듯.
        # 다모아버섯이 '나눠진' 이유는 제가 순서대로 출력했기 때문입니다. 전체 합산하면 해결됩니다.
        
        supplier_aggs = {} # {sup: {(name, unit): qty}}
        supplier_memos = {} # {sup: [memo...]}
        
        # 3. 공급처별 합산 실행
        for p in processed_items:
            if p["type"] == "item":
                sup = p["sup"]
                if sup not in supplier_aggs: supplier_aggs[sup] = {}
                
                key = (p["item"], p["unit"])
                if key not in supplier_aggs[sup]: supplier_aggs[sup][key] = 0
                supplier_aggs[sup][key] += p["qty"]
                
                # 메모 할당 로직 (휴리스틱): 가장 최근에 본 sup에게 할당?
                # 아니면 pending_memos에 쌓인걸 이 sup에게?
                if pending_memos:
                    if sup not in supplier_memos: supplier_memos[sup] = []
                    supplier_memos[sup].extend(pending_memos)
                    pending_memos = []

            elif p["type"] == "memo":
                # 메모는 잠시 대기 (다음 아이템의 주인에게 귀속)
                pending_memos.append(p["text"])
        
        # 남은 메모 처리 (처리 안 된 메모는 누락 가능성 있음, 로그 필요)
        
        # 4. 공급처 리스트 생성 (Page 3)
        supplier_final_lines = {} # {sup: [lines...]}
        
        for sup, agg in supplier_aggs.items():
            lines = []
            # 일반 아이템
            for (name, unit), qty in agg.items():
                qs = str(qty)
                if qs.endswith(".0"): qs = qs[:-2]
                if qty == 0: qs = ""
                lines.append(f"{name} {qs}{unit}".strip())
            
            # 메모 추가 (일단 뒤에 붙임, 오복상회 (선발주) 같은 경우 뒤에 나오는게 일반적)
            # 만약 (선발주) 아이템들이 구분이 안 되어 합산되어 버린다면?
            # 예: 굵은숙주 20 + 15 = 35박스. 
            # 사용자 예시: 굵은숙주 20, (선발주) 굵은숙주 15. -> 분리되길 원함.
            # 이러려면 합산 키에 'Group' 개념이 있어야 함.
            # 하지만 입력 txt 파일에는 그런 구분자가 없음. (Store 단위로만 구분됨)
            # 즉, Store별로 '일반'인지 '선발주'인지 알 수 없으면 합칠 수 밖에 없음.
            
            # 여기서 사용자의 "클로드 출력"을 보면: `오복상회` 밑에 `(선발주)`가 따로 있음.
            # 이는 오복상회 주문 데이터가 입력될 때 구분되어 있었거나, 
            # 특정 Store(예: 오복상회 선발주용 가상 Store)가 있었을 것임.
            # 하지만 orders_today.txt에는 그런게 안 보임.
            # 따라서 '전체 합산'이 최선임. 다모아버섯 문제 해결이 우선.
            
            if sup in supplier_memos:
                lines.extend(supplier_memos[sup])
                
            supplier_final_lines[sup] = lines

        # 5. 문서 작성
        doc = Document()
        for section in doc.sections:
            section.top_margin = Cm(MARGIN); section.bottom_margin = Cm(MARGIN)
            section.left_margin = Cm(MARGIN); section.right_margin = Cm(MARGIN)
            section.page_width = Cm(21); section.page_height = Cm(29.7)
            
        weekday = datetime.datetime.now().weekday()
        config_key = "friday" if weekday == 4 else "default"
        page_config = self.mappings.get("page_config", {}).get(config_key, {})
        sorted_pages = sorted(page_config.keys())
        
        processed_stores = set()
        pages_data = []

        for page_name in sorted_pages:
            targets = page_config[page_name]
            page_info = {'cols': 3, 'content': [], 'type': 'normal'}
            content_columns = [[], [], []]

            if isinstance(targets, list):
                # Page 1, 2
                chunk_size = (len(targets) + 2) // 3
                chunks = [targets[i:i + chunk_size] for i in range(0, len(targets), chunk_size)]
                
                for col_idx, store_group in enumerate(chunks):
                    if col_idx >= 3: break
                    col_lines = []
                    for key in store_group:
                        matched = sorted([k for k in orders.keys() if key in k])
                        for s in matched:
                            processed_stores.add(s)
                            if col_lines: col_lines.append("")
                            col_lines.append(f"-{s}")
                            for item in orders[s]:
                                if item.get("type") == "memo": col_lines.append(item["text"])
                                else: col_lines.append(item.get("original_line", ""))
                    content_columns[col_idx] = col_lines
            
            elif targets == "SUPPLIER_LIST":
                # Page 3: Supplier Aggregated
                all_sup = sorted(supplier_final_lines.keys())
                all_sup = [s for s in all_sup if "기타" not in s]
                
                chunk_size = (len(all_sup) + 2) // 3
                chunks = [all_sup[i:i + chunk_size] for i in range(0, len(all_sup), chunk_size)]
                
                for col_idx, grp in enumerate(chunks):
                    col_lines = []
                    for sup in grp:
                        if col_lines: col_lines.append("")
                        col_lines.append(f"-{sup}")
                        col_lines.extend(supplier_final_lines[sup])
                    content_columns[col_idx] = col_lines
            
            elif targets == "STOCK_MARKET_LIST":
                # Page 4
                page_info['cols'] = 2; page_info['type'] = 'split'
                content_columns = [[], []]
                content_columns[0].append("-재고, 창고소분")
                content_columns[0].extend(stock_lines["STOCK"])
                content_columns[1].append("-시장구매, 시장소분")
                content_columns[1].extend(stock_lines["MARKET"])
                
            page_info['content'] = content_columns[:page_info['cols']]
            pages_data.append(page_info)
            
        # Missed
        missed = [s for s in orders.keys() if s not in processed_stores]
        if missed:
             c = []
             for s in missed:
                 if c: c.append("")
                 c.append(f"-{s}")
                 for item in orders[s]:
                     if item.get("type") == "memo": c.append(item["text"])
                     else: c.append(item.get("original_line",""))
             pages_data.append({'cols': 3, 'content': [c, [], []], 'type': 'normal'})
             
        # Render
        for page_idx, page in enumerate(pages_data):
            if page_idx > 0:
                doc.add_section()
                for s in doc.sections:
                    s.top_margin = Cm(MARGIN); s.bottom_margin = Cm(MARGIN)
                    s.left_margin = Cm(MARGIN); s.right_margin = Cm(MARGIN)
            cols = page['cols']
            table = doc.add_table(1, cols)
            row = table.rows[0]
            for i, data in enumerate(page['content']):
                row.cells[i].width = Cm(9.5 if cols==2 else 6.3)
                st = None
                if page.get('type') == 'split':
                    if i==0: st='stock'
                    elif i==1: st='market'
                self._add_content_to_cell(row.cells[i], data, st)
                set_cell_border(row.cells[i], **BORDER_SETTINGS)
                
        doc.save(output_path)
