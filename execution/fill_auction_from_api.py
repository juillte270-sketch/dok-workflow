"""
가락시장 공공데이터 API에서 경매가를 조회하여
도크발주관리데이터.xlsx 단가 시트에 최고가/평균가를 자동 채우는 스크립트.

API: 품목별등급별가격(도매시장법인거래)
- XML endpoint: dataXmlOpen.do (dataid=data68)
- 한글 파라미터는 UTF-8 URLEncoding 사용

Usage:
  python execution/fill_auction_from_api.py \
    --master "G:/내 드라이브/1. 도크_주문 명세서/0. 매입단가_자료/도크발주관리데이터.xlsx" \
    --date 2026-02-23
"""
import argparse
import os
import re
import time
import urllib.parse
from datetime import datetime, timedelta

import openpyxl
import requests

# API 기본 설정
API_BASE = "http://www.garak.co.kr/homepage/publicdata/dataXmlOpen.do"
API_ID = "6335"
API_PASSWD = "tfe7c1p4!!"
API_DATAID = "data68"

# 단가 시트 품목명 → API 검색 품목명 (일반 매칭용, ITEM_OVERRIDES에 없는 품목만 사용)
ITEM_SEARCH_MAP = {
    '청경채': '청경채',
    '쑥갓': '쑥갓',
    '알배기배추': '알배기배추',
    '근대': '근대',
    '깻잎(1KG)': '깻잎',
    '깻잎(4KG,큰잎)': '깻잎순',
    '치커리': '치커리',
    '케일': '케일',
    '팽이': '팽이버섯',
    '새송이': '새송이버섯',
    '맛느타리': '맛느타리버섯',
    '꽃느타리': '꽃느타리버섯',
    '양송이': '양송이',
    '청양고추': '청양고추',
    '양파': '양파',
    '깐양파': '양파',
    '감자': '감자',
    '당근 수입': '당근',
    '흙대파': '대파',
    '깐쪽파': '깐쪽파',
    '부추(일반)': '부추',
    '돌미나리': '돌미나리',
    '미나리': '미나리',
    '취청오이': '취청오이',
    '배추': '배추',
    '깐양배추45': '양배추',
    '양배추45': '양배추',
    '귤': '감귤',
    '양상추(일반)': '양상추',
    '양상추(국산/수입)': '양상추',
    '양상추': '양상추',
    '라디치오': None,       # 수입, 경매 데이터 없음
    '간마늘': None,          # 가공품, 경매 데이터 없음
    '아욱': '아욱',
    '깐마늘 대서': '깐마늘',
    '깐대파': '대파',
    '애호박': '애호박',
    '꽈리고추': '꽈리고추',
    '생표고 수입': '생표고 수입',
    '로메인(일반)': '로메인',
    '통로메인': '로메인',
    '레몬': None,            # 수입과일, 경매 데이터 없음
    '홍고추': '홍고추',
    '무': '무',
    '팽이버섯': '팽이버섯',
    '시금치': '시금치',
}

# 특수 매칭 규칙 — 품목별 필터/변환이 필요한 경우 (ITEM_SEARCH_MAP 대신 우선 적용)
# search: API 검색어
# filter_item_exact: 결과 중 품목명이 정확히 일치하는 것만
# filter_item_contains: 결과 중 품목명에 해당 문자열 포함된 것만
# target_grade: 매칭할 등급
# target_unit_kg: 특정 kg 단위 매칭
# unit_contains: 거래단위 문자열에 해당 키워드 포함
# price_factor: 가격 변환 비율 (예: 0.5=50%, 0.1=1/10, 0.05=1/20)
# note: 셀 서식에 텍스트 표시 (예: '흙' → 셀에 "13,900(흙)" 으로 표시, 숫자 값은 유지)
ITEM_OVERRIDES = {
    # 깻잎(1KG): 깻잎 상 100속 가격의 50% (50속 가격)
    '깻잎(1KG)': {
        'search': '깻잎',
        'filter_item_exact': '깻잎',
        'target_grade': '상',
        'unit_contains': '속',
        'price_factor': 0.5,
    },
    # 깻잎(4KG,큰잎): 깻잎순 상 4KG상자 가격 그대로
    '깻잎(4KG,큰잎)': {
        'search': '깻잎순',
        'target_grade': '상',
        'target_unit_kg': 4.0,
        'price_factor': 1.0,
    },
    # 깐쪽파: 깐쪽파 특 10KG상자 가격의 1/10 (1K 단가)
    '깐쪽파': {
        'search': '깐쪽파',
        'target_grade': '특',
        'target_unit_kg': 10.0,
        'price_factor': 0.1,
    },
    # 미나리: 미나리 상 4KG 가격의 1/4 (1K 단가) — 돌미나리 제외
    '미나리': {
        'search': '미나리',
        'filter_item_exact': '미나리',
        'target_grade': '상',
        'target_unit_kg': 4.0,
        'price_factor': 0.25,
    },
    # 깐양배추45: 양배추 특 8KG그물망 가격 + "흙" 표시
    '깐양배추45': {
        'search': '양배추',
        'filter_item_exact': '양배추',
        'target_grade': '특',
        'unit_contains': '그물망',
        'price_factor': 1.0,
        'note': '흙',
    },
    # 양상추: 양상추(일반) 상 12개 가격 — 수입 제외
    '양상추(일반)': {
        'search': '양상추',
        'filter_item_contains': '일반',
        'target_grade': '상',
        'unit_contains': '개',
        'price_factor': 1.0,
    },
    '양상추(국산/수입)': {
        'search': '양상추',
        'filter_item_contains': '일반',
        'target_grade': '상',
        'unit_contains': '개',
        'price_factor': 1.0,
    },
    '양상추': {
        'search': '양상추',
        'filter_item_contains': '일반',
        'target_grade': '상',
        'unit_contains': '개',
        'price_factor': 1.0,
    },
    # 배추: 배추 특 10Kg그물망 — '배추' 검색 시 양배추/알배기 등도 나오므로 정확 필터 필수
    '배추': {
        'search': '배추',
        'filter_item_exact': '배추',
        'target_grade': '특',
        'unit_contains': '그물망',
        'price_factor': 1.0,
    },
    # 토마토 완숙 — 단가시트명 그대로 (CJ: 완숙토마토 특 5kg상자)
    '토마토 완숙': {
        'search': '토마토',
        'filter_item_contains': '완숙',
        'target_grade': '특',
        'target_unit_kg': 5.0,
        'price_factor': 1.0,
    },
    # 브로콜리 국산 — 수입과 구분 필요 (filter_item_exact)
    '브로콜리 국산': {
        'search': '브로콜리',
        'filter_item_exact': '브로콜리 국산',
        'target_grade': '상',
        'target_unit_kg': 8.0,
        'price_factor': 1.0,
    },
    # 귤: 감귤 5kg 상 등급
    '귤': {
        'search': '감귤',
        'filter_item_exact': '감귤',
        'target_grade': '상',
        'target_unit_kg': 5.0,
        'price_factor': 1.0,
    },
    # 깐마늘 대서: API는 20KG만 있으므로 시트 단위에 맞게 자동 변환
    # use_sheet_grade=True → 시트의 등급(상or대→상, 보통or중→보통)으로 API 매칭
    # auto_convert_from_kg=20 → API 20KG 가격을 시트 단위(1K, 20K)로 자동 변환
    '깐마늘 대서': {
        'search': '깐마늘',
        'target_unit_kg': 20.0,
        'use_sheet_grade': True,
        'auto_convert_from_kg': 20.0,
    },
    # 빨간양배추 국산: '양배추'로 검색 → '빨간양배추 국산' 필터 (API에 '적양배추'는 없음)
    '빨간양배추 국산': {
        'search': '양배추',
        'filter_item_exact': '빨간양배추 국산',
        'target_grade': '상',
        'target_unit_kg': 8.0,
        'price_factor': 1.0,
    },
}

# 품목별 셀 표시 — override/일반 매칭 모두에 적용
# 가격 셀에 숫자 서식 #,##0"(흙)" 적용 → "13,900(흙)" 표시, 값은 숫자 유지
ITEM_NOTES = {
    '깐양배추45': '흙',
    '깐양파': '흙',
    '깐대파': '흙',
}


def get_prev_business_day(dt):
    """일요일을 제외한 전일 계산"""
    prev = dt - timedelta(days=1)
    while prev.weekday() == 6:  # Sunday = 6
        prev -= timedelta(days=1)
    return prev


def parse_price(val):
    """콤마 포함 가격 문자열을 정수로 변환"""
    if not val:
        return None
    s = str(val).replace(',', '').strip()
    if not s or s == '-':
        return None
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return None


def parse_unit_kg(unit_str):
    """API 거래단위에서 kg 수치 추출. '15 kg' → 15, '500 g' → 0.5"""
    if not unit_str:
        return None
    s = str(unit_str).lower().strip()
    m = re.search(r'([\d.]+)\s*kg', s)
    if m:
        return float(m.group(1))
    m = re.search(r'([\d.]+)\s*g', s)
    if m:
        return float(m.group(1)) / 1000
    return None


def parse_danga_unit_kg(unit_str):
    """단가 시트 단위에서 kg 수치 추출. '15K' → 15, '1K' → 1"""
    if not unit_str:
        return None
    s = str(unit_str).upper().strip()
    m = re.search(r'([\d.]+)\s*K', s)
    if m:
        return float(m.group(1))
    return None


def fetch_auction_data(search_item, s_date, s_date_p, s_date_p7):
    """API에서 품목 검색 → [{item, grade, unit_str, unit_kg, max_price, avg_price}, ...]"""
    encoded_item = urllib.parse.quote(search_item, encoding='utf-8')
    url = (
        f"{API_BASE}?id={API_ID}&passwd={API_PASSWD}&dataid={API_DATAID}"
        f"&pagesize=100&pageidx=1&portal.templet=false"
        f"&s_date={s_date}&s_date_p={s_date_p}&s_date_p7={s_date_p7}"
        f"&p_pos_gubun=1&s_pum_nm=2&s_pummok={encoded_item}"
    )

    try:
        resp = requests.get(url, timeout=20)
        resp.raise_for_status()
    except requests.RequestException as e:
        print(f"    API request error: {e}")
        return []

    text = resp.text

    # XML이 malformed될 수 있으므로 regex로 파싱
    results = []
    items = re.findall(r'<list>(.*?)</list>', text, re.DOTALL)

    for item_xml in items:
        def extract(tag):
            m = re.search(rf'<{tag}><!\[CDATA\[(.*?)\]\]></{tag}>', item_xml)
            return m.group(1).strip() if m else ''

        pum_nm = extract('PUM_NM')
        g_name = extract('G_NAME')
        u_name = extract('U_NAME')
        ma_p = parse_price(extract('MA_P'))
        av_p = parse_price(extract('AV_P'))
        unit_kg = parse_unit_kg(u_name)

        if ma_p is not None and av_p is not None:
            results.append({
                'item': pum_nm,
                'grade': g_name,
                'unit_str': u_name,
                'unit_kg': unit_kg,
                'max_price': ma_p,
                'avg_price': av_p,
            })

    return results


def find_override_match(auction_results, override):
    """Override 규칙에 따라 경매 결과에서 필터 → 등급 매칭 → 가격 변환"""
    entries = list(auction_results)

    # 1. 품목명 필터
    if 'filter_item_exact' in override:
        exact = override['filter_item_exact']
        entries = [e for e in entries if e['item'] == exact]
    if 'filter_item_contains' in override:
        kw = override['filter_item_contains']
        entries = [e for e in entries if kw in e['item']]

    # 2. 단위 필터 (kg 수치)
    if 'target_unit_kg' in override:
        target_kg = override['target_unit_kg']
        kg_entries = [e for e in entries if e['unit_kg'] == target_kg]
        if kg_entries:
            entries = kg_entries

    # 3. 단위 필터 (문자열 키워드)
    if 'unit_contains' in override:
        uc = override['unit_contains']
        uc_entries = [e for e in entries if uc in str(e['unit_str'])]
        if uc_entries:
            entries = uc_entries

    if not entries:
        return None, None, None, None

    # 4. 등급 매칭
    target_grade = override.get('target_grade')
    matched = None
    if target_grade:
        for e in entries:
            if e['grade'] == target_grade:
                matched = e
                break
    if not matched:
        grade_prio = {e['grade']: e for e in entries}
        for g in ['특', '상', '보통', '하']:
            if g in grade_prio:
                matched = grade_prio[g]
                break
    if not matched:
        matched = entries[0]

    # 5. 가격 변환
    factor = override.get('price_factor', 1.0)
    max_p = int(round(matched['max_price'] * factor))
    avg_p = int(round(matched['avg_price'] * factor))

    return max_p, avg_p, matched['grade'], matched['unit_str']


def find_best_match(auction_results, target_grade, target_kg):
    """경매 결과에서 등급+단위 매칭하여 최적 가격 반환 (일반 품목용)"""
    if not auction_results:
        return None, None, None, None

    entries = auction_results

    # 1. 단위 매칭: target_kg에 가장 가까운 단위 그룹 선택
    if target_kg:
        kg_entries = [e for e in entries if e['unit_kg'] is not None]
        if kg_entries:
            available_kgs = set(e['unit_kg'] for e in kg_entries)
            closest_kg = min(available_kgs, key=lambda x: abs(x - target_kg))
            entries = [e for e in kg_entries if e['unit_kg'] == closest_kg]

    # 2. 등급 매칭
    grade_map = {e['grade']: e for e in entries}

    # 등급 정규화
    normalized_grade = target_grade
    if target_grade in ['상or대', '상OR대', '상Or대']:
        normalized_grade = '상'
    elif target_grade in ['보통or중', '보통OR중', '보통Or중']:
        normalized_grade = '보통'

    # target_grade 우선
    if normalized_grade in grade_map:
        e = grade_map[normalized_grade]
        return e['max_price'], e['avg_price'], e['grade'], e['unit_str']

    # 없으면 등급 우선순위대로
    for g in ['특', '상', '보통', '하']:
        if g in grade_map:
            e = grade_map[g]
            return e['max_price'], e['avg_price'], e['grade'], e['unit_str']

    return None, None, None, None


def fill_auction_prices(master_file, target_date, auction_date=None):
    """메인 함수: 단가 시트에서 대상 행을 찾고 API로 경매가 채우기

    Args:
        master_file: 도크발주관리데이터.xlsx 경로
        target_date: 단가 시트 날짜 (YYYY-MM-DD) - 발주일
        auction_date: 경매 API 조회 날짜 (YYYY-MM-DD) - 기본값: 발주일+1일
    """
    dt = datetime.strptime(target_date, "%Y-%m-%d")

    # 경매 API 날짜 = 발주일 + 1일 (경매는 당일~익일새벽, API는 익일로 등록)
    if auction_date:
        auction_dt = datetime.strptime(auction_date, "%Y-%m-%d")
    else:
        auction_dt = dt + timedelta(days=1)

    s_date = auction_dt.strftime("%Y%m%d")
    prev_day = get_prev_business_day(auction_dt)
    s_date_p = prev_day.strftime("%Y%m%d")
    s_date_p7 = (auction_dt - timedelta(days=7)).strftime("%Y%m%d")

    print(f"단가시트 날짜: {target_date} / 경매 API 날짜: {auction_dt.strftime('%Y-%m-%d')}")
    print(f"API params: s_date={s_date}, s_date_p={s_date_p}, s_date_p7={s_date_p7}")

    print(f"\nOpening master file: {master_file}")
    wb = openpyxl.load_workbook(master_file)
    ws = wb["단가"]

    # 대상 행 수집 (첫 번째 연속 블록만)
    target_rows = []
    in_block = False
    last_match_row = 0
    for row_idx in range(2, ws.max_row + 1):
        date_cell = ws.cell(row=row_idx, column=2).value
        item_cell = ws.cell(row=row_idx, column=3).value

        is_match = False
        if date_cell:
            if str(date_cell).startswith(target_date):
                is_match = True
            elif hasattr(date_cell, 'strftime') and date_cell.strftime("%Y-%m-%d") == target_date:
                is_match = True

        if is_match and item_cell:
            # 첫 번째 블록이 끝난 후 다시 매칭되면 중단 (gap > 2행)
            if in_block and (row_idx - last_match_row) > 2:
                break
            in_block = True
            last_match_row = row_idx

            item_name = str(item_cell).strip()
            category = str(ws.cell(row=row_idx, column=1).value or '')
            grade = str(ws.cell(row=row_idx, column=4).value or '').strip()
            unit = str(ws.cell(row=row_idx, column=5).value or '').strip()
            target_rows.append((row_idx, item_name, category, grade, unit))

    print(f"Found {len(target_rows)} items for {target_date}")

    # API 결과 캐시
    api_cache = {}
    updates = 0
    skipped = 0
    no_match = []

    for row_idx, item_name, category, grade, unit in target_rows:
        # 소분/재고 행은 스킵
        if '소분' in category:
            skipped += 1
            continue

        note = None
        max_price = avg_price = matched_grade = matched_unit = None
        factor_info = ""

        # === Override 우선 확인 ===
        if item_name in ITEM_OVERRIDES:
            override = dict(ITEM_OVERRIDES[item_name])  # copy (동적 수정용)
            search_key = override['search']

            # 시트 등급을 그대로 사용 (상or대→상, 보통or중→보통 정규화)
            if override.get('use_sheet_grade'):
                norm_grade = grade
                if '상' in grade and 'or' in grade.lower():
                    norm_grade = '상'
                elif '보통' in grade and 'or' in grade.lower():
                    norm_grade = '보통'
                override['target_grade'] = norm_grade

            # API 단위(예:20KG)에서 시트 단위(예:1K)로 자동 변환
            if 'auto_convert_from_kg' in override:
                source_kg = override['auto_convert_from_kg']
                target_kg = parse_danga_unit_kg(unit)
                if target_kg and source_kg:
                    override['price_factor'] = target_kg / source_kg

            if search_key not in api_cache:
                print(f"  API call: '{search_key}'...", end=" ", flush=True)
                results = fetch_auction_data(search_key, s_date, s_date_p, s_date_p7)
                api_cache[search_key] = results
                print(f"{len(results)} results")
                time.sleep(0.3)

            auction_results = api_cache[search_key]
            max_price, avg_price, matched_grade, matched_unit = find_override_match(
                auction_results, override
            )
            note = override.get('note')
            pf = override.get('price_factor', 1.0)
            if pf != 1.0:
                factor_info = f" x{pf}"

        # === 일반 매칭 ===
        else:
            if item_name not in ITEM_SEARCH_MAP:
                no_match.append((row_idx, item_name, "No mapping defined"))
                continue

            search_key = ITEM_SEARCH_MAP[item_name]
            if search_key is None:
                no_match.append((row_idx, item_name, "Excluded (no auction data)"))
                continue

            if search_key not in api_cache:
                print(f"  API call: '{search_key}'...", end=" ", flush=True)
                results = fetch_auction_data(search_key, s_date, s_date_p, s_date_p7)
                api_cache[search_key] = results
                print(f"{len(results)} results")
                time.sleep(0.3)

            auction_results = api_cache[search_key]
            target_kg = parse_danga_unit_kg(unit)

            max_price, avg_price, matched_grade, matched_unit = find_best_match(
                auction_results, grade, target_kg
            )

        # === 결과 기록 ===
        # override의 note 또는 ITEM_NOTES에서 표시 텍스트 결정
        display_note = note or ITEM_NOTES.get(item_name)

        if max_price and avg_price:
            cell_f = ws.cell(row=row_idx, column=6)
            cell_g = ws.cell(row=row_idx, column=7)

            if not isinstance(cell_f, openpyxl.cell.cell.MergedCell):
                cell_f.value = max_price
                if display_note:
                    cell_f.number_format = f'#,##0"({display_note})"'
            if not isinstance(cell_g, openpyxl.cell.cell.MergedCell):
                cell_g.value = avg_price
                if display_note:
                    cell_g.number_format = f'#,##0"({display_note})"'

            note_str = f" ({display_note})" if display_note else ""
            print(f"  R{row_idx}: {item_name}({grade} {unit}) -> {matched_grade} {matched_unit}{factor_info} "
                  f"최고={max_price:,} 평균={avg_price:,}{note_str}")
            updates += 1
        else:
            no_match.append((row_idx, item_name, f"No auction match for '{search_key}'"))

    print(f"\n=== Results ===")
    print(f"Updated: {updates}")
    print(f"Skipped (소분/재고): {skipped}")
    print(f"No match: {len(no_match)}")
    for r, name, reason in no_match:
        print(f"  R{r}: {name} - {reason}")

    if updates > 0:
        # 저장 (잠금 시 재시도)
        saved = False
        for attempt in range(5):
            try:
                print(f"\nSaving to {master_file}... (attempt {attempt + 1})")
                wb.save(master_file)
                print("Done.")
                saved = True
                break
            except PermissionError:
                if attempt < 4:
                    print(f"  File locked, retrying in 3s...")
                    time.sleep(3)
                else:
                    # 임시 파일로 저장
                    import tempfile
                    tmp = os.path.join(tempfile.gettempdir(), "도크발주관리데이터_auction.xlsx")
                    wb.save(tmp)
                    print(f"  Saved to temp: {tmp}")
                    print(f"  Please close the file and copy manually.")
                    saved = True
    else:
        print("\nNo updates made.")

    wb.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="가락시장 API로 경매가(최고가/평균가)를 단가 시트에 자동 채우기"
    )
    parser.add_argument("--master", required=True, help="도크발주관리데이터.xlsx 경로")
    parser.add_argument("--date", required=True, help="단가시트 발주일 (YYYY-MM-DD)")
    parser.add_argument("--auction-date", default=None,
                        help="경매 API 조회일 (YYYY-MM-DD). 기본값: 발주일+1일")
    args = parser.parse_args()

    import sys
    try:
        fill_auction_prices(args.master, args.date, args.auction_date)
    except SystemExit:
        raise
    except Exception:
        import traceback; traceback.print_exc()
        sys.exit(1)

    from _notify import notify
    notify("stage5_auction", date_str=args.date)
