from docx import Document
import os

def test_library():
    print("--- Testing Library ---")
    try:
        doc = Document()
        doc.add_paragraph("Hello world")
        print("Created document object successfully")
        doc.save("test_gen.docx")
        print("Saved test_gen.docx")
        
        doc2 = Document("test_gen.docx")
        print(f"Read back test_gen.docx: {doc2.paragraphs[0].text}")
        print("Library works fine.\n")
    except Exception as e:
        print(f"Library test failed: {e}")

def analyze_docx(file_path):
    print(f"Checking file: {file_path}")
    if not os.path.exists(file_path):
        print("Error: File does not exist!")
        return

    try:
        doc = Document(file_path)
        print(f"=== Document Analysis: {file_path} ===\n")
        
        print("--- Paragraphs ---")
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                print(f"[P{i}] Style: {para.style.name} | Text: {para.text[:50]}...")
        
        print("\n--- Tables ---")
        for i, table in enumerate(doc.tables):
            print(f"[Table {i}] Rows: {len(table.rows)} Cols: {len(table.columns)}")
            try:
                for r_idx, row in enumerate(table.rows[:3]): # 3줄만
                    row_cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    print(f"  R{r_idx}: {row_cells}")
            except Exception as e:
                print(f"  Error reading rows: {e}")
                
        print("\n=== End Analysis ===")
        
    except Exception as e:
        print(f"Error analyzing docx: {e}")

if __name__ == "__main__":
    test_library()
    
    # 다른 파일 시도
    path = r"c:\Users\DoKH_D\OneDrive\Desktop\일일 직납 현황\일일_직납_현황_260127_1.docx"
    analyze_docx(path)
