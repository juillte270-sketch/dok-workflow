from docx import Document
import os

def analyze_docx(file_path):
    print(f"Checking file: {file_path}")
    if not os.path.exists(file_path):
        print("Error: File does not exist!")
        return

    try:
        doc = Document(file_path)
        print(f"=== Document Analysis: {file_path} ===\n")
        
        print("--- Paragraphs ---")
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                print(f"[P{i}] Style: {para.style.name} | Text: {para.text[:50]}...")
        
        print("\n--- Tables ---")
        for i, table in enumerate(doc.tables):
            print(f"[Table {i}] Rows: {len(table.rows)} Cols: {len(table.columns)}")
            
            # 첫 몇 줄만 출력해서 구조 파악
            try:
                for r_idx, row in enumerate(table.rows[:5]):
                    row_cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    print(f"  R{r_idx}: {row_cells}")
            except Exception as e:
                print(f"  Error reading rows: {e}")
                
        print("\n=== End Analysis ===")
        
    except Exception as e:
        print(f"Error analyzing docx: {e}")

if __name__ == "__main__":
    # 경로를 역슬래시로 변경하고 raw string 사용
    path = r"c:\Users\DoKH_D\OneDrive\Desktop\일일 직납 현황\일일_직납_현황_260128.docx"
    analyze_docx(path)
