import openpyxl
from datetime import datetime
import os
import argparse

# 마스터 파일 경로
DEFAULT_MASTER_FILE = r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터.xlsx"
SHEET_NAME = "단가"

def clean_price_sheet(target_date, master_file):
    print(f"Opening {master_file}...")
    try:
        wb = openpyxl.load_workbook(master_file)
    except Exception as e:
        print(f"Error opening file: {e}")
        return

    if SHEET_NAME not in wb.sheetnames:
        print(f"Sheet '{SHEET_NAME}' not found.")
        return

    ws = wb[SHEET_NAME]
    
    # Target date parsing
    if isinstance(target_date, str):
        target_date = datetime.strptime(target_date, "%Y-%m-%d").date()
    
    print(f"Scanning for data with date: {target_date}...")
    
    rows_to_delete = []
    # Scan B column (Date)
    for row in range(ws.max_row, 1, -1):
        cell_val = ws.cell(row=row, column=2).value
        # Check if it matches target date
        if isinstance(cell_val, datetime):
            if cell_val.date() == target_date:
                rows_to_delete.append(row)
        elif isinstance(cell_val,str):
            # Sometimes date might be string?
            try:
                dt = datetime.strptime(cell_val[:10], "%Y-%m-%d").date()
                if dt == target_date:
                     rows_to_delete.append(row)
            except:
                pass

    if rows_to_delete:
        print(f"Found {len(rows_to_delete)} rows to delete.")
        # Delete rows
        for r in rows_to_delete:
            ws.delete_rows(r, 1)
        
        print("Saving workbook...")
        wb.save(master_file)
        print("Done.")
    else:
        print("No rows found for that date.")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True, help="YYYY-MM-DD")
    parser.add_argument("--master", default=DEFAULT_MASTER_FILE)
    args = parser.parse_args()
    
    clean_price_sheet(args.date, args.master)
