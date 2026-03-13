"""
배송관리 시트 생성: Firestore 배송 데이터 → 마스터 엑셀 "배송관리" 시트 기록.

매장별 상세 행 + 기사별 합계 행으로 기록.

Usage:
    python execution/generate_delivery_report.py --date 20260313
    python execution/generate_delivery_report.py --date 20260313 --dry-run

Requires: firebase-admin, openpyxl
"""
import argparse
import math
import os
import re
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

# ─── Paths ────────────────────────────────────────────────────
PROJECT_ROOT = Path(__file__).resolve().parent.parent
MASTER_FILE = Path(
    r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터.xlsx"
)
SHEET_NAME = "배송관리"
KST = timezone(timedelta(hours=9))

# Driver UID → Name mapping
DRIVER_MAP = {
    "5oTVDPG2Wqgirrro08otaXcLvhe2": "배승길",
    "sqpKNlwam7W2pxra8jNolzLaGxF3": "이형기",
    "1jU70ijZqIMp6iqCg6z37BaSvxG2": "박용구",
}

# Headers (품목수/박스수/총무게는 총(분) 뒤)
HEADERS = [
    "날짜", "기사", "#", "매장",
    "출발", "도착", "완료",
    "이동(분)", "입고(분)", "총(분)",
    "품목수", "박스수", "총무게",
    "km",
]

# ─── 매장 좌표 (배송 거리 계산용) ──────────────────────────────
GARAK_MARKET = (37.503860, 127.113830)  # 가락시장 (배송 출발점)

STORE_COORDS = [
    ("육회", 37.495400, 127.123900),
    ("봄날", 37.492947, 127.124287),
    ("고른햇살", 37.590068, 127.030305),
    ("소유", 37.538620, 127.056875),
    ("선데이", 37.526549, 127.036813),
    ("이너프유", 37.502248, 126.789145),
    ("브럭시", 37.513742, 127.104247),
    ("압구정", 37.527328, 127.037997),
    ("하남", 37.545445, 127.224057),
    ("강남점", 37.500078, 127.027516),
    ("마곡", 37.559229, 126.835888),
    ("구월", 37.445662, 126.703161),
    ("주안", 37.461984, 126.671774),
    ("분당", 37.349277, 127.108402),
    ("평택", 37.033484, 127.013226),
    ("동탄", 37.200359, 127.095576),
    ("도봉", 37.648923, 127.035000),
    ("광교", 37.287526, 127.055817),
    ("역삼", 37.499555, 127.047845),
    ("송파", 37.493820, 127.122433),
    ("동원", 37.197509, 126.991205),
]


def _get_store_coord(store_name: str) -> tuple[float, float] | None:
    name_norm = store_name.replace(" ", "")
    for keyword, lat, lng in STORE_COORDS:
        if keyword in name_norm:
            return (lat, lng)
    return None


def _haversine_km(c1: tuple[float, float], c2: tuple[float, float]) -> float:
    """두 좌표 사이 직선 거리 (km)."""
    R = 6371.0
    lat1, lon1 = math.radians(c1[0]), math.radians(c1[1])
    lat2, lon2 = math.radians(c2[0]), math.radians(c2[1])
    dlat, dlon = lat2 - lat1, lon2 - lon1
    a = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return R * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))


# ─── Items 파싱 ───────────────────────────────────────────────

def _parse_items_info(items: list[dict]) -> tuple[int, int, float]:
    """items → (품목수, 박스수, 총무게kg).

    qty가 "kg" 포함 → 총무게에 합산, 아니면 → 박스수에 합산.
    """
    item_count = len(items)
    box_count = 0
    total_kg = 0.0

    for item in items:
        qty_str = str(item.get("qty", "1"))
        num_m = re.match(r"(\d+(?:\.\d+)?)", qty_str)
        num = float(num_m.group(1)) if num_m else 1

        if "kg" in qty_str.lower():
            total_kg += num
        else:
            box_count += int(num)

    return item_count, box_count, total_kg


# ─── Firebase Init ────────────────────────────────────────────

def _get_firestore_client():
    """Firebase Admin SDK 초기화 + Firestore 클라이언트 반환."""
    import firebase_admin
    from firebase_admin import credentials, firestore

    if not firebase_admin._apps:
        service_key_path = os.environ.get("FIREBASE_SERVICE_ACCOUNT_KEY")
        if not service_key_path or not Path(service_key_path).exists():
            service_key_path = str(PROJECT_ROOT / "serviceAccountKey.json")
        if not Path(service_key_path).exists():
            print("[ERROR] serviceAccountKey.json을 찾을 수 없습니다.")
            sys.exit(1)
        cred = credentials.Certificate(service_key_path)
        firebase_admin.initialize_app(cred)

    return firestore.client()


# ─── Firestore Data Fetch ─────────────────────────────────────

def fetch_deliveries(db, date_compact: str) -> list[dict]:
    """deliveryDays/{date}/deliveries 전체 조회."""
    ref = db.collection("deliveryDays").document(date_compact).collection("deliveries")
    results = []
    for doc in ref.stream():
        d = doc.to_dict()
        d["id"] = doc.id
        results.append(d)
    return results


def fetch_vehicle_km(db, driver_uid: str) -> float | None:
    """vehicles/{driverId}.totalKm 조회."""
    doc = db.collection("vehicles").document(driver_uid).get()
    if doc.exists:
        return doc.to_dict().get("totalKm")
    return None


# ─── Timestamp Helpers ────────────────────────────────────────

def _ts_to_dt(ts) -> datetime | None:
    if ts is None:
        return None
    if hasattr(ts, "seconds"):
        dt_utc = datetime.fromtimestamp(ts.seconds + ts.nanoseconds / 1e9, tz=timezone.utc)
        return dt_utc.astimezone(KST)
    if isinstance(ts, datetime):
        if ts.tzinfo is None:
            return ts.replace(tzinfo=KST)
        return ts.astimezone(KST)
    return None


def _fmt_time(dt: datetime | None) -> str:
    if dt is None:
        return ""
    return dt.strftime("%H:%M")


def _minutes_between(start: datetime | None, end: datetime | None) -> int:
    if start is None or end is None:
        return 0
    delta = (end - start).total_seconds() / 60.0
    return round(max(delta, 0.0))


# ─── Build Detail Rows ───────────────────────────────────────

def build_report_rows(deliveries: list[dict], db) -> list[dict]:
    """매장별 상세 행 + 기사별 합계 행 생성.

    Returns list of row dicts (type='detail' or type='summary').
    """
    # Group by driver UID
    by_driver: dict[str, list[dict]] = {}
    for d in deliveries:
        uid = d.get("assignedTo")
        if not uid:
            continue
        by_driver.setdefault(uid, []).append(d)

    all_rows = []

    for uid, dds in sorted(by_driver.items(), key=lambda x: DRIVER_MAP.get(x[0], x[0])):
        driver_name = DRIVER_MAP.get(uid)
        if not driver_name:
            for dd in dds:
                if dd.get("driverName"):
                    driver_name = dd["driverName"]
                    break
            if not driver_name:
                driver_name = uid[:8]

        # Sort by order
        dds.sort(key=lambda x: x.get("order", 999))

        # Per-store detail rows
        sum_items = 0
        sum_boxes = 0
        sum_kg = 0.0
        sum_transit = 0
        sum_unload = 0
        sum_total = 0
        sum_km = 0.0
        start_times = []
        end_times = []

        prev_coord = GARAK_MARKET  # 첫 출발: 가락시장

        for seq, d in enumerate(dds, 1):
            started = _ts_to_dt(d.get("startedAt"))
            arrived = _ts_to_dt(d.get("arrivedAt"))
            done = _ts_to_dt(d.get("completedAt"))

            transit = _minutes_between(started, arrived)
            unload = _minutes_between(arrived, done)
            total = _minutes_between(started, done)

            items = d.get("items", [])
            items_count, boxes, kg = _parse_items_info(items)
            sum_items += items_count
            sum_boxes += boxes
            sum_kg += kg
            sum_transit += transit
            sum_unload += unload
            sum_total += total

            if started:
                start_times.append(started)
            if done:
                end_times.append(done)

            # km: 이전 위치 → 현재 매장 (직선 거리)
            store_name_raw = d.get("storeName", d.get("id", "?"))
            cur_coord = _get_store_coord(store_name_raw)
            segment_km = None
            if cur_coord and prev_coord:
                segment_km = round(_haversine_km(prev_coord, cur_coord), 1)
                sum_km += segment_km
                prev_coord = cur_coord

            status = d.get("status", "")
            store_name = store_name_raw
            if status != "delivered":
                store_name += f" ({status})"

            all_rows.append({
                "type": "detail",
                "driver": driver_name,
                "seq": seq,
                "store": store_name,
                "start": _fmt_time(started),
                "arrive": _fmt_time(arrived),
                "complete": _fmt_time(done),
                "transit_min": transit,
                "unload_min": unload,
                "total_min": total,
                "items": items_count,
                "boxes": boxes,
                "kg": kg,
                "km": segment_km,
            })

        # Summary row
        earliest = min(start_times) if start_times else None
        latest = max(end_times) if end_times else None
        work_min = _minutes_between(earliest, latest)
        completed_count = sum(1 for d in dds if d.get("status") == "delivered")

        all_rows.append({
            "type": "summary",
            "driver": driver_name,
            "seq": "",
            "store": f"합계 ({completed_count}/{len(dds)})",
            "start": _fmt_time(earliest),
            "arrive": "",
            "complete": _fmt_time(latest),
            "transit_min": sum_transit,
            "unload_min": sum_unload,
            "total_min": work_min,
            "items": sum_items,
            "boxes": sum_boxes,
            "kg": round(sum_kg, 1),
            "km": round(sum_km, 1),
        })

    return all_rows


# ─── Excel Write ──────────────────────────────────────────────

NUM_COLS = len(HEADERS)  # 14


def write_to_excel(date_str: str, rows: list[dict], dry_run: bool = False):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    if not MASTER_FILE.exists():
        print(f"[ERROR] 마스터 파일을 찾을 수 없습니다: {MASTER_FILE}")
        sys.exit(1)

    wb = openpyxl.load_workbook(str(MASTER_FILE))
    formatted_date = f"{date_str[:4]}-{date_str[4:6]}-{date_str[6:8]}"

    center = Alignment(horizontal="center", vertical="center")

    # Get or create sheet
    new_sheet = SHEET_NAME not in wb.sheetnames
    if new_sheet:
        ws = wb.create_sheet(SHEET_NAME)
        print(f"  [INFO] '{SHEET_NAME}' 시트 새로 생성")
    else:
        ws = wb[SHEET_NAME]

    # Write headers if new or row 1 is empty
    if new_sheet or ws.cell(row=1, column=1).value is None:
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_text = Font(bold=True, size=10, color="FFFFFF")
        for col_idx, header in enumerate(HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_text
            cell.fill = header_fill
            cell.alignment = center
        ws.freeze_panes = "A2"
        # Column widths (14 cols)
        widths = [12, 8, 4, 20, 7, 7, 7, 8, 8, 8, 6, 6, 8, 7]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Check for duplicate date
    for row in range(2, ws.max_row + 1):
        cell_val = ws.cell(row=row, column=1).value
        if cell_val is None:
            continue
        cell_str = str(cell_val).replace("/", "-").strip()
        if cell_str == formatted_date or cell_str == date_str:
            print(f"  [SKIP] {formatted_date} 데이터가 이미 존재합니다 (행 {row})")
            wb.close()
            return

    # Find first empty row
    first_empty = ws.max_row + 1
    if first_empty <= 1:
        first_empty = 2

    if dry_run:
        print(f"\n  [DRY-RUN] 행 {first_empty}부터 {len(rows)}행 추가 예정")
        wb.close()
        return

    # Styles
    summary_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    summary_font = Font(bold=True, size=10)
    thin_border = Border(bottom=Side(style="thin", color="C0C0C0"))
    summary_border = Border(
        top=Side(style="medium", color="4472C4"),
        bottom=Side(style="medium", color="4472C4"),
    )

    # Write rows
    for i, r in enumerate(rows):
        row_num = first_empty + i
        vals = [
            formatted_date,                                     # 1 날짜
            r["driver"],                                        # 2 기사
            r["seq"],                                           # 3 #
            r["store"],                                         # 4 매장
            r["start"],                                         # 5 출발
            r["arrive"],                                        # 6 도착
            r["complete"],                                      # 7 완료
            r["transit_min"] if r["transit_min"] else "",        # 8 이동(분)
            r["unload_min"] if r["unload_min"] else "",         # 9 입고(분)
            r["total_min"] if r["total_min"] else "",           # 10 총(분)
            r["items"] if r["items"] else "",                   # 11 품목수
            r["boxes"] if r.get("boxes") else "",               # 12 박스수
            r["kg"] if r.get("kg") else "",                     # 13 총무게
            r["km"] if r.get("km") is not None else "",         # 14 km
        ]

        for col_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.alignment = center

        # Apply styles
        if r["type"] == "summary":
            for col in range(1, NUM_COLS + 1):
                cell = ws.cell(row=row_num, column=col)
                cell.fill = summary_fill
                cell.font = summary_font
                cell.border = summary_border
                cell.alignment = center
        else:
            for col in range(1, NUM_COLS + 1):
                cell = ws.cell(row=row_num, column=col)
                cell.border = thin_border
                cell.alignment = center

    wb.save(str(MASTER_FILE))
    wb.close()
    last_row = first_empty + len(rows) - 1
    print(f"  [OK] {SHEET_NAME} 시트에 {len(rows)}행 저장 (행 {first_empty}~{last_row})")


# ─── CLI ──────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="배송관리 시트 생성 (Firestore → Excel)")
    parser.add_argument("--date", help="날짜 (YYYYMMDD, 기본: 오늘)")
    parser.add_argument("--master", help="마스터 파일 경로 (기본: G드라이브)")
    parser.add_argument("--dry-run", action="store_true", help="미리보기 (엑셀 저장 안 함)")
    args = parser.parse_args()

    global MASTER_FILE
    if args.master:
        MASTER_FILE = Path(args.master)

    if args.date:
        date_str = args.date.replace("-", "")
    else:
        date_str = datetime.now(KST).strftime("%Y%m%d")

    if len(date_str) != 8:
        print(f"[ERROR] 날짜 형식 오류: {args.date}")
        sys.exit(1)

    formatted_date = f"{date_str[:4]}-{date_str[4:6]}-{date_str[6:8]}"
    print(f"[배송관리 리포트] {formatted_date}")

    # 1. Fetch from Firestore
    print("\n  Firestore에서 배송 데이터 조회 중...")
    db = _get_firestore_client()
    deliveries = fetch_deliveries(db, date_str)

    if not deliveries:
        print(f"  [WARN] {formatted_date} 배송 데이터가 없습니다.")
        sys.exit(0)

    print(f"  총 {len(deliveries)}건 배송 조회됨")

    # 2. Build detail + summary rows
    rows = build_report_rows(deliveries, db)

    if not rows:
        print(f"  [WARN] 배정된 기사가 없습니다.")
        sys.exit(0)

    # 3. Print summary
    print()
    current_driver = None
    for r in rows:
        if r["driver"] != current_driver:
            current_driver = r["driver"]
            print(f"  [{current_driver}]")
        km_str = f"{r['km']:.1f}km" if r.get("km") else ""
        if r["type"] == "detail":
            print(f"    #{r['seq']} {r['store']:<18s} {r['items']}개/{r.get('boxes',0)}박스  "
                  f"출발{r['start']:>5s} 도착{r['arrive']:>5s} 완료{r['complete']:>5s}  "
                  f"이동{r['transit_min']:>3d}분 입고{r['unload_min']:>3d}분 총{r['total_min']:>3d}분  {km_str}")
        else:
            print(f"    ── {r['store']}  "
                  f"{r['start']:>5s}~{r['complete']:>5s}  "
                  f"이동{r['transit_min']:>3d}분 입고{r['unload_min']:>3d}분 총{r['total_min']:>3d}분  {km_str}")

    # 4. Write to Excel
    print(f"\n  마스터 파일: {MASTER_FILE}")
    write_to_excel(date_str, rows, dry_run=args.dry_run)

    print(f"\n완료.")


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception:
        import traceback
        traceback.print_exc()
        sys.exit(1)
