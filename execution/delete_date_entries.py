import openpyxl
from datetime import datetime
import os

# 마스터 파일 경로
MASTER_FILE = r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터.xlsx"
sheet_name = "발주"

# 삭제할 날짜
target_date = datetime(2026, 2, 3).date()

print(f"Opening {MASTER_FILE}...")
wb = openpyxl.load_workbook(MASTER_FILE)
ws = wb[sheet_name]

# 2026-02-03 데이터 찾기
rows_to_delete = []
max_row = ws.max_row

print(f"Scanning {max_row} rows for 2026-02-03 data...")
for r in range(2, max_row + 1):
    date_val = ws.cell(row=r, column=4).value  # D열 (날짜)
    if isinstance(date_val, datetime):
        if date_val.date() == target_date:
            rows_to_delete.append(r)

if rows_to_delete:
    print(f"Found {len(rows_to_delete)} rows with 2026-02-03 data")
    print(f"Rows: {rows_to_delete[0]} to {rows_to_delete[-1]}")
    
    # 역순으로 삭제 (인덱스 변경 방지)
    for r in reversed(rows_to_delete):
        ws.delete_rows(r, 1)
    
    wb.save(MASTER_FILE)
    print(f"✓ Deleted {len(rows_to_delete)} rows")
else:
    print("No 2026-02-03 data found")
