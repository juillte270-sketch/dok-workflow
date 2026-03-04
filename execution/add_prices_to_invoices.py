"""
최종 거래명세서 생성 스크립트 (간단 버전)

기존 generate_invoices.py로 생성된 거래명세서에 단가 정보만 추가합니다.

사용법:
    python execution/add_prices_to_invoices.py --date "2026-02-03"
"""

import os
import glob
import openpyxl
from openpyxl.styles import Font
from datetime import datetime, timedelta
import argparse

# 프로젝트 루트
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# 마스터 파일 경로
MASTER_FILE = r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터.xlsx"
MASTER_SHEET = "발주"

def load_master_prices(target_date):
    """마스터 파일에서 판매가 정보 로드"""
    print(f"\n📊 마스터 파일에서 판매가 로드 중...")
    
    wb = openpyxl.load_workbook(MASTER_FILE, data_only=True)
    ws = wb[MASTER_SHEET]
    
    # 날짜별, 매장별, 품목별 판매가 딕셔너리
    prices = {}
    
    print(f"   총 행 수: {ws.max_row:,}개")
    
    for row in range(2, ws.max_row + 1):
        if row % 5000 == 0:
            print(f"   처리 중: {row:,}/{ws.max_row:,} 행...")
        
        매장명 = ws.cell(row, 1).value
        날짜 = ws.cell(row, 4).value
        품목 = ws.cell(row, 5).value
        판매가 = ws.cell(row, 11).value
        총판매금액 = ws.cell(row, 12).value
        
        if not all([매장명, 날짜, 품목]):
            continue
        
        # 날짜 변환
        if isinstance(날짜, datetime):
            날짜_str = 날짜.strftime("%Y-%m-%d")
        else:
            continue
        
        # 키 생성: "날짜|매장명|품목"
        key = f"{날짜_str}|{매장명}|{품목}"
        
        # 안전한 float 변환
        def safe_float(value):
            if value is None or value == '' or value == '-':
                return 0
            try:
                return float(value)
            except:
                return 0
        
        prices[key] = {
            '판매가': safe_float(판매가),
            '총판매금액': safe_float(총판매금액)
        }
    
    wb.close()
    print(f"✅ 판매가 정보 로드 완료: {len(prices):,}개")
    return prices

def add_prices_to_invoice(invoice_path, prices, target_date, store_name):
    """거래명세서에 판매가 추가"""
    print(f"\n💰 {store_name} - 판매가 입력 중...")
    print(f"   파일: {os.path.basename(invoice_path)}")
    
    wb = openpyxl.load_workbook(invoice_path)
    sheet = wb.active
    
    # 품목 시작 행 (16행부터)
    start_row = 16
    
    # 단가 열과 공급가액 열 찾기 (15행이 헤더)
    단가_col = None
    공급가액_col = None
    
    for col in range(1, 20):
        header = sheet.cell(15, col).value  # 15행이 헤더
        if header and '단가' in str(header):
            단가_col = col
        if header and ('공급가액' in str(header) or '금액' in str(header)):
            공급가액_col = col
    
    if not 단가_col or not 공급가액_col:
        print(f"   ⚠️  단가 또는 공급가액 열을 찾을 수 없습니다.")
        wb.close()
        return None
    
    print(f"   ✓ 단가 열: {chr(64+단가_col)} ({단가_col})")
    print(f"   ✓ 공급가액 열: {chr(64+공급가액_col)} ({공급가액_col})")
    
    # 품목별 판매가 입력
    current_row = start_row
    daily_total = 0
    items_processed = 0
    
    while current_row < 100:
        품목_cell = sheet.cell(current_row, 2)  # B열: 품목
        
        if not 품목_cell.value or str(품목_cell.value).strip() == '':
            break
        
        품목 = str(품목_cell.value).strip()
        
        # "합계" 등이 나오면 중단
        if '합계' in 품목 or '계' == 품목:
            break
        
        # 마스터 파일에서 판매가 찾기
        # 파일명의 언더스코어를 공백으로 변환 (마스터 파일은 공백 사용)
        store_name_for_key = store_name.replace('_', ' ')
        key = f"{target_date.strftime('%Y-%m-%d')}|{store_name_for_key}|{품목}"
        
        if key in prices:
            판매가 = prices[key]['판매가']
            총판매금액 = prices[key]['총판매금액']
            
            # 단가 입력
            sheet.cell(current_row, 단가_col).value = 판매가
            sheet.cell(current_row, 단가_col).number_format = '#,##0'
            
            # 공급가액 입력
            sheet.cell(current_row, 공급가액_col).value = 총판매금액
            sheet.cell(current_row, 공급가액_col).number_format = '#,##0'
            
            daily_total += 총판매금액
            items_processed += 1
        else:
            print(f"   ⚠️  판매가 없음: {품목}")
        
        current_row += 1
    
    print(f"   ✓ 처리 완료: {items_processed}개 품목")
    print(f"   ✓ 당일금액: {daily_total:,}원")
    
    # 당월금액 계산
    monthly_amount = calculate_monthly_amount(prices, store_name, target_date)
    total_amount = monthly_amount + daily_total
    
    # 금액 입력 (하드코딩된 위치 - 템플릿에 따라 조정 필요)
    # 상단 합계금액 (예: G8)
    try:
        sheet['G8'] = daily_total
        sheet['G8'].number_format = '#,##0'
        print(f"   ✓ 상단 합계금액: {daily_total:,}원")
    except:
        pass
    
    # 하단 금액 정보 (예: G41~G45)
    try:
        sheet['G41'] = monthly_amount
        sheet['G41'].number_format = '#,##0'
        print(f"   ✓ 당월금액: {monthly_amount:,}원")
        
        sheet['G42'] = daily_total
        sheet['G42'].number_format = '#,##0'
        print(f"   ✓ 당일금액: {daily_total:,}원")
        
        sheet['G43'] = total_amount
        sheet['G43'].number_format = '#,##0'
        print(f"   ✓ 총사용액: {total_amount:,}원")
        
        # 입금액 확인 (있으면 사용, 없으면 0)
        입금액 = sheet['G44'].value or 0
        
        # 하단 합계금액
        final_balance = total_amount - 입금액
        sheet['G45'] = final_balance
        sheet['G45'].number_format = '#,##0'
        sheet['G45'].font = Font(bold=True)
        print(f"   ✓ 하단 합계금액: {final_balance:,}원")
    except Exception as e:
        print(f"   ⚠️  금액 입력 오류: {e}")
    
    # 저장
    output_filename = f"{store_name}_최종명세서_{target_date.strftime('%Y%m%d')}.xlsx"
    output_dir = os.path.join(PROJECT_ROOT, "data", "outputs", f"final_invoices_{target_date.strftime('%Y%m%d')}")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, output_filename)
    
    wb.save(output_path)
    print(f"   ✅ 저장 완료: {output_filename}")
    
    return output_path

def calculate_monthly_amount(prices, store_name, target_date):
    """당월금액 계산 (1일 ~ 전일)"""
    if target_date.day == 1:
        return 0
    
    month_start = target_date.replace(day=1)
    day_before = target_date - timedelta(days=1)
    
    total = 0
    
    # 파일명의 언더스코어를 공백으로 변환
    store_name_for_key = store_name.replace('_', ' ')
    
    # 해당 기간의 모든 날짜 확인
    current_date = month_start
    while current_date <= day_before:
        date_str = current_date.strftime("%Y-%m-%d")
        
        # 해당 날짜의 모든 품목 합계
        for key, value in prices.items():
            if key.startswith(f"{date_str}|{store_name_for_key}|"):
                total += value['총판매금액']
        
        current_date += timedelta(days=1)
    
    return total

def convert_to_pdf(excel_path):
    """Excel을 PDF로 변환"""
    try:
        import win32com.client
        
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        
        wb = excel.Workbooks.Open(os.path.abspath(excel_path))
        pdf_path = excel_path.replace('.xlsx', '.pdf')
        
        # 1페이지로 설정
        ws = wb.Worksheets(1)
        ws.PageSetup.Zoom = False
        ws.PageSetup.FitToPagesWide = 1
        ws.PageSetup.FitToPagesTall = 1
        
        wb.ExportAsFixedFormat(0, pdf_path)
        wb.Close(False)
        excel.Quit()
        
        print(f"   ✅ PDF 저장: {os.path.basename(pdf_path)}")
        return pdf_path
        
    except Exception as e:
        print(f"   ⚠️  PDF 변환 실패: {e}")
        return None

def main():
    parser = argparse.ArgumentParser(description="거래명세서에 판매가 추가")
    parser.add_argument("--date", help="대상 날짜 (YYYY-MM-DD)", required=True)
    args = parser.parse_args()
    
    try:
        target_date = datetime.strptime(args.date, "%Y-%m-%d").date()
    except ValueError:
        print("❌ 오류: 날짜 형식이 올바르지 않습니다. YYYY-MM-DD 형식으로 입력해주세요.")
        return
    
    print(f"""
╔══════════════════════════════════════════════════════════╗
║         거래명세서 판매가 추가 시스템                     ║
╚══════════════════════════════════════════════════════════╝

📅 대상 날짜: {args.date}
""")
    
    # 1. 마스터 파일에서 판매가 로드
    prices = load_master_prices(target_date)
    
    # 2. 생성된 거래명세서 찾기
    date_str = target_date.strftime("%Y%m%d")
    invoice_dir = os.path.join(PROJECT_ROOT, "data", "outputs", f"invoices_{date_str}")
    
    if not os.path.exists(invoice_dir):
        print(f"\n❌ 오류: 거래명세서 폴더가 없습니다: {invoice_dir}")
        print(f"먼저 다음 명령을 실행하세요:")
        print(f"  python execution/generate_invoices.py --date {args.date}")
        return
    
    # 3. 각 거래명세서에 판매가 추가
    invoice_files = glob.glob(os.path.join(invoice_dir, "*.xlsx"))
    invoice_files = [f for f in invoice_files if not os.path.basename(f).startswith('~$')]
    
    print(f"\n🏪 처리 대상: {len(invoice_files)}개 거래명세서")
    
    success_count = 0
    fail_count = 0
    
    for invoice_file in invoice_files:
        # 파일명에서 매장명 추출 (예: "20260203_샤브야키_도봉점_거래명세서.xlsx")
        basename = os.path.basename(invoice_file)
        parts = basename.replace('.xlsx', '').split('_')
        
        if len(parts) >= 3:
            store_name = '_'.join(parts[1:-1])  # 날짜와 "거래명세서" 제외
        else:
            print(f"\n⚠️  건너뜀: 파일명 형식 오류 - {basename}")
            fail_count += 1
            continue
        
        print(f"\n{'='*60}")
        print(f"📋 {store_name}")
        print(f"{'='*60}")
        
        # 판매가 추가
        output_path = add_prices_to_invoice(invoice_file, prices, target_date, store_name)
        
        if output_path:
            # PDF 변환
            convert_to_pdf(output_path)
            success_count += 1
        else:
            fail_count += 1
    
    # 요약
    print(f"\n{'='*60}")
    print(f"🎉 처리 완료!")
    print(f"{'='*60}")
    print(f"   성공: {success_count}개")
    print(f"   실패: {fail_count}개")

if __name__ == "__main__":
    main()
