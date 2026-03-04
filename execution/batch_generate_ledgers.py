"""
Batch Transaction Ledger Generator
Generates transaction ledgers for multiple stores for February 1-6, 2026
"""

import pandas as pd
import openpyxl
import os
import shutil
from datetime import datetime
import zipfile

# Store mapping: folder name -> display name in master data
STORE_MAPPING = {
    "0. 샤브야키 광교점": "샤브야키_광교점",
    "1. 샤브야키 분당점": "샤브야키_분당점",
    "2. 샤브야키 평택점": "샤브야키_평택점",
    "3. 샤브야키 주안점": "샤브야키_주안점",
    "4. 샤브야키 동탄점": "샤브야키_동탄점",
    "5. 부엉이 산장": ["부엉이산장_강남점", "부엉이산장_구월점", "부엉이산장_마곡점"],
    "6. 고른햇살": "고른햇살",
    "7. 오레노이키루미치": ["오레노이키루미치_압구정점", "오레노이키루미치_하남점"],
    "8. 성동 과일 클래스 소유": "소유프루트",  # 소유프루트 = 소유
    "9. 가락 과일 클래스 봄날": "봄날",
    "11. 브럭시": "브럭시",
    "12. 선데이버거클럽": "선데이버거클럽",
    "13. 육회꽃필무렵 송파점": "육회꽃필무렵_송파점",
    "14. 파라이 카이카이": ["파라이_송파점", "파라이_역삼점"],
    "0. 이너프유": "이너프유",
}

# Configuration
MASTER_DATA_PATH = r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터.xlsx"
TRANSACTION_BASE_PATH = r"G:\내 드라이브\1. 도크_주문 명세서\1. 거래명세서"
OUTPUT_ZIP_PATH = r"c:\Users\DoKH_D\OneDrive\Desktop\안티그래비티 프로젝트 (3)\data\outputs\거래원장_2026년2월_1-6일.zip"
START_DATE = "2026-02-01"
END_DATE = "2026-02-06"


def find_template_file(store_folder):
    """Find the template file for a given store folder"""
    template_pattern = "(거래원장 템플릿)"
    for file in os.listdir(store_folder):
        if template_pattern in file and file.endswith(".xlsx"):
            return os.path.join(store_folder, file)
    return None


def load_master_data(store_name, start_date, end_date):
    """Load and aggregate master data for a specific store and date range"""
    df = pd.read_excel(MASTER_DATA_PATH, sheet_name='발주')
    
    # Convert dates
    df['기준일'] = pd.to_datetime(df['기준일'])
    start_dt = pd.to_datetime(start_date)
    end_dt = pd.to_datetime(end_date)
    
    # Filter by store and date range
    mask = (df['매장'] == store_name) & (df['기준일'] >= start_dt) & (df['기준일'] <= end_dt)
    filtered_df = df[mask]
    
    # Aggregate by date
    daily_totals = filtered_df.groupby('기준일')['총 판매금액'].sum().reset_index()
    daily_totals = daily_totals.sort_values('기준일')
    
    return daily_totals


def generate_ledger(store_folder, store_names, start_date, end_date):
    """Generate transaction ledger for a store"""
    print(f"\n{'='*60}")
    print(f"Processing: {store_folder}")
    print(f"{'='*60}")
    
    # Find template
    template_path = find_template_file(os.path.join(TRANSACTION_BASE_PATH, store_folder))
    if not template_path:
        print(f"  ❌ Template not found for {store_folder}")
        return None
    
    print(f"  Template: {os.path.basename(template_path)}")
    
    # Load data for all store names (some stores have multiple locations)
    if isinstance(store_names, list):
        all_daily_totals = pd.DataFrame()
        for store_name in store_names:
            daily_totals = load_master_data(store_name, start_date, end_date)
            all_daily_totals = pd.concat([all_daily_totals, daily_totals])
        # Re-aggregate if multiple locations
        daily_totals = all_daily_totals.groupby('기준일')['총 판매금액'].sum().reset_index()
        daily_totals = daily_totals.sort_values('기준일')
    else:
        daily_totals = load_master_data(store_names, start_date, end_date)
    
    if daily_totals.empty:
        print(f"  ⚠️  No data found for {store_folder}")
        return None
    
    print(f"  Data rows: {len(daily_totals)}")
    
    # Create output folder
    start_dt = pd.to_datetime(start_date)
    folder_name = f"{start_dt.year}년 {start_dt.month}월"
    output_folder = os.path.join(TRANSACTION_BASE_PATH, store_folder, folder_name)
    os.makedirs(output_folder, exist_ok=True)
    
    # Create output file
    file_name = f"{start_dt.year}{start_dt.month:02d}_{store_folder.split('. ')[1]}_월말정산_거래원장.xlsx"
    output_path = os.path.join(output_folder, file_name)
    
    # Copy template
    shutil.copy(template_path, output_path)
    
    # Update ledger
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    
    # Clear existing data (rows 17-43)
    for r in range(17, 44):
        ws.cell(row=r, column=2).value = None  # Date
        ws.cell(row=r, column=13).value = None  # Amount
    
    # Write data
    start_row = 17
    total_sum = 0
    for idx, row in daily_totals.iterrows():
        current_row = start_row + idx
        ws.cell(row=current_row, column=2).value = row['기준일']
        ws.cell(row=current_row, column=2).number_format = 'yyyy-mm-dd'
        ws.cell(row=current_row, column=13).value = row['총 판매금액']
        ws.cell(row=current_row, column=13).number_format = '#,##0'
        total_sum += row['총 판매금액']
    
    # Update title (B5)
    month_str = f"{start_dt.month:02d}월"
    title_text = f"{store_folder.split('. ')[1]}\\n{month_str} 월말 정산 내역"
    ws.cell(row=5, column=2).value = title_text
    
    # Update date range (D7)
    end_dt = pd.to_datetime(end_date)
    date_range_text = f"{start_dt.year}년 {start_dt.month:02d}월 {start_dt.day:02d}일 ~ {end_dt.month:02d}월 {end_dt.day:02d}일"
    ws.cell(row=7, column=4).value = date_range_text
    
    # Update bottom summary fields
    ws.cell(row=47, column=6).value = total_sum
    ws.cell(row=47, column=6).number_format = '#,##0'
    ws.cell(row=49, column=6).value = total_sum
    ws.cell(row=49, column=6).number_format = '#,##0'
    
    wb.save(output_path)
    print(f"  ✅ Saved: {file_name}")
    print(f"  Total: {total_sum:,}원")
    
    return output_path


def create_zip_file(generated_files):
    """Create a zip file with all generated ledgers"""
    os.makedirs(os.path.dirname(OUTPUT_ZIP_PATH), exist_ok=True)
    
    with zipfile.ZipFile(OUTPUT_ZIP_PATH, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for file_path in generated_files:
            if file_path:
                arcname = os.path.basename(file_path)
                zipf.write(file_path, arcname)
    
    print(f"\n{'='*60}")
    print(f"✅ Zip file created: {OUTPUT_ZIP_PATH}")
    print(f"   Total files: {len([f for f in generated_files if f])}")
    print(f"{'='*60}")


def main():
    """Main execution function"""
    print("="*60)
    print("Batch Transaction Ledger Generator")
    print(f"Period: {START_DATE} ~ {END_DATE}")
    print("="*60)
    
    generated_files = []
    
    for store_folder, store_names in STORE_MAPPING.items():
        try:
            output_path = generate_ledger(store_folder, store_names, START_DATE, END_DATE)
            generated_files.append(output_path)
        except Exception as e:
            print(f"  ❌ Error processing {store_folder}: {e}")
            import traceback
            traceback.print_exc()
    
    # Create zip file
    create_zip_file(generated_files)
    
    print("\n✅ Batch processing complete!")


if __name__ == "__main__":
    main()
