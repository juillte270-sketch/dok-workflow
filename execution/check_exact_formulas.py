import openpyxl
import os

master_file_path = r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터-AI접근용.xlsx"
sheet_name = "발주"

def check_exact_formulas(rows=[19000, 19408]):
    if not os.path.exists(master_file_path):
        print(f"Error: Master file not found at {master_file_path}")
        return

    try:
        wb = openpyxl.load_workbook(master_file_path, data_only=False)
        ws = wb[sheet_name]
        
        with open(".tmp/exact_formula_check.txt", "w", encoding="utf-8") as f:
            for row in rows:
                f.write(f"--- Row {row} ---\n")
                col_m = ws.cell(row=row, column=13) # M: 마진율
                col_n = ws.cell(row=row, column=14) # N: 이익률
                f.write(f"Col M (Margin): {col_m.value}\n")
                f.write(f"Col N (Profit): {col_n.value}\n")
                f.write("\n")
                
        print("✓ Exact formula check written to .tmp/exact_formula_check.txt")

    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    check_exact_formulas()
