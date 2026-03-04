import os
import glob
import shutil
import openpyxl
import argparse
from openpyxl.styles import Border, Side, Font
from openpyxl.cell.cell import MergedCell
from datetime import datetime
from collections import defaultdict
import re
import unicodedata

# ---------------------------------------------------------
# 환경 설정
# ---------------------------------------------------------
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER_FILE = r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터.xlsx"
DRIVE_ROOT = r"G:\내 드라이브\1. 도크_주문 명세서\1. 거래명세서"

# ---------------------------------------------------------
# 보조 함수
# ---------------------------------------------------------

def normalize_key(text):
    """텍스트 정규화 (공백 제거 및 NFC)"""
    if not text: return ""
    return unicodedata.normalize('NFC', str(text)).replace(' ', '').strip()

def find_drive_folder(root_path, store_name):
    """거래처 명칭으로 드라이브 내 폴더 경로 찾기"""
    folder_map = {
        '샤브야키 분당점': '1. 샤브야키 분당점',
        '샤브야키 평택점': '2. 샤브야키 평택점',
        '샤브야키 주안점': '3. 샤브야키 주안점',
        '샤브야키 동탄점': '4. 샤브야키 동탄점',
        '샤브야키 도봉점': '0. 샤브야키 도봉점',
        '샤브야키 광교점': '0. 샤브야키 광교점',
        '부엉이산장 강남점': r'5. 부엉이 산장\1. 부엉이산장 강남점',
        '부엉이산장 지오다노': r'5. 부엉이 산장\1. 부엉이산장 강남점',
        '부엉이산장 구월점': r'5. 부엉이 산장\2. 부엉이산장 구월점',
        '부엉이산장 마곡점': r'5. 부엉이 산장\3. 부엉이산장 마곡점',
        '오레노이키루미치 하남점': r'7. 오레노이키루미치\오레노이키루미치 하남점',
        '오레노이키루미치 압구정점': r'7. 오레노이키루미치\오레노이키루미치 압구정점',
        '고른햇살': '6. 고른햇살',
        '브럭시': '11. 브럭시',
        '선데이버거클럽': '12. 선데이버거클럽',
        '육회꽃필무렵': '13. 육회꽃필무렵 송파점',
        '소유 프루트': '8. 성동 과일 클래스 소유',
        '소유프루트': '8. 성동 과일 클래스 소유',
        '봄날': '9. 가락 과일 클래스 봄날',
        '이너프유': '0. 이너프유',
        '파라이 송파점': r'14. 파라이 카이카이\파라이 송파점',
        '파라이 역삼점': r'14. 파라이 카이카이\파라이 역삼점',
        '다비다': '10. 다비다웨딩홀뷔페 연제점',
    }
    
    # 1. Exact Match
    if store_name in folder_map:
        return folder_map[store_name]
    
    # 2. Key matching
    for key, folder in folder_map.items():
        if normalize_key(key) == normalize_key(store_name):
            return folder
            
    # 3. Partial Match
    matched_keys = sorted([k for k in folder_map.keys() if k in store_name or store_name in k], key=len, reverse=True)
    if matched_keys:
        return folder_map[matched_keys[0]]
        
    return store_name 

def find_template(store_name, target_date):
    """가장 적절한 템플릿 파일을 찾음 (AI참고용 우선)"""
    drive_folder = find_drive_folder(DRIVE_ROOT, store_name)
    store_path = os.path.join(DRIVE_ROOT, drive_folder)
    
    if not os.path.exists(store_path):
        return None
    
    # 1. "AI참고용" 검색
    for root, dirs, files in os.walk(store_path):
        for f in files:
            if ("AI참고용" in f or "AI" in f) and f.endswith(".xlsx") and not f.startswith("~$"):
                return os.path.join(root, f)
                
    # 2. 최신 월 폴더에서 검색
    month_folder = f"{target_date.year}년 {target_date.month}월"
    month_path = os.path.join(store_path, month_folder)
    if os.path.exists(month_path):
        files = [f for f in os.listdir(month_path) if f.endswith(".xlsx") and not f.startswith("~$")]
        if files:
            files.sort(key=lambda x: os.path.getmtime(os.path.join(month_path, x)), reverse=True)
            return os.path.join(month_path, files[0])
            
    return None

def load_order_data(data_path, year, month):
    """마스터 데이터 로드 (정규화 키로 매장명 통합)"""
    print(f"[INFO] Loading Master Data... ({year}-{month})")
    wb = openpyxl.load_workbook(data_path, data_only=True)
    ws = wb['발주']

    store_totals = defaultdict(lambda: defaultdict(float))
    order_items = defaultdict(lambda: defaultdict(dict))
    # 정규화키 → 원본 매장명 매핑 (첫 번째 등장 이름 사용)
    norm_to_original = {}

    count = 0
    for row in range(2, ws.max_row + 1):
        store = ws.cell(row=row, column=1).value
        date = ws.cell(row=row, column=4).value
        item = ws.cell(row=row, column=5).value
        qty = ws.cell(row=row, column=8).value
        price = ws.cell(row=row, column=11).value

        if not all([store, date, item]): continue

        if not isinstance(date, datetime):
            try: date = datetime.strptime(str(date)[:10], "%Y-%m-%d")
            except: continue

        if date.year == year and date.month == month:
            s_key = normalize_key(store)
            day = date.day

            # 정규화 키로 첫 번째 원본 이름 기억
            if s_key not in norm_to_original:
                norm_to_original[s_key] = store

            try: q = float(qty) if qty else 0
            except: q = 0
            try: p = float(price) if price else 0
            except: p = 0

            # 정규화 키로 통합 저장 (매장명 변형 문제 해결)
            store_totals[s_key][day] += q * p
            order_items[s_key][day][item] = (q, p)
            count += 1

    wb.close()
    print(f"[INFO] Data loaded: {count} items")
    return store_totals, order_items, norm_to_original

def safe_set_value(ws, row, col, value, bold=False):
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for mr in ws.merged_cells.ranges:
            if cell.coordinate in mr:
                cell = ws.cell(mr.min_row, mr.min_col)
                break
    cell.value = value
    if isinstance(value, (int, float)):
        cell.number_format = '#,##0'
    if bold:
        cell.font = Font(bold=True)

def safe_set_border(ws, row, col, border):
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        try: cell.border = border
        except: pass

def convert_to_pdf_win32(excel_path):
    import win32com.client
    try:
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        abs_path = os.path.abspath(excel_path)
        wb = excel.Workbooks.Open(abs_path)
        pdf_path = abs_path.replace('.xlsx', '.pdf')
        wb.ExportAsFixedFormat(0, pdf_path)
        wb.Close(False)
        excel.Quit()
        return pdf_path
    except Exception as e:
        print(f"   [WARN] PDF failed: {e}")
        return None

def process_statement(template_path, store_name_master, day, output_path, store_totals, order_items):
    shutil.copy(template_path, output_path)
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    
    # 전화번호 고정 (010-3275-0188)
    for row in range(1, 15):
        for col in range(1, 15):
            cell = ws.cell(row=row, column=col)
            if cell.value and '010' in str(cell.value):
                safe_set_value(ws, row, col, '010-3275-0188')

    # 행 찾기
    monthly_row = 41
    for row in range(35, 60):
        if ws.cell(row, 2).value and '당월' in str(ws.cell(row, 2).value):
            monthly_row = row
            break
            
    sum_row = monthly_row - 2
    for r in range(monthly_row - 1, 20, -1):
        if ws.cell(r, 2).value and ('합계' in str(ws.cell(r, 2).value) or '계' == str(ws.cell(r, 2).value).strip()):
            sum_row = r
            break
            
    # 데이터
    orders = order_items[store_name_master].get(day, {})
    prev_total = sum(store_totals[store_name_master].get(d, 0) for d in range(1, day))
    
    # 품목 초기화
    for row in range(16, sum_row):
        for col in [2, 5, 10, 11]: safe_set_value(ws, row, col, None)
            
    # 품목 입력
    total_amount = 0
    current_row = 16
    for item_name, (qty, price) in orders.items():
        if current_row >= sum_row: break
        safe_set_value(ws, current_row, 2, item_name)
        safe_set_value(ws, current_row, 5, qty)
        safe_set_value(ws, current_row, 10, price)
        safe_set_value(ws, current_row, 11, qty * price)
        total_amount += qty * price
        current_row += 1
        
    # 합계 및 하단
    safe_set_value(ws, sum_row, 11, total_amount)
    safe_set_value(ws, 8, 4, total_amount)
    safe_set_value(ws, monthly_row, 4, int(prev_total))
    safe_set_value(ws, monthly_row + 1, 4, int(total_amount))
    safe_set_value(ws, monthly_row + 2, 4, int(prev_total + total_amount))
    safe_set_value(ws, monthly_row + 4, 4, int(prev_total + total_amount))
    
    # 빈 행 숨김
    no_border = Border(left=Side(style=None), right=Side(style=None), top=Side(style=None), bottom=Side(style=None))
    for row in range(current_row, sum_row):
        ws.row_dimensions[row].hidden = True
        for col in range(2, 15): safe_set_border(ws, row, col, no_border)

    wb.save(output_path)
    wb.close()
    return total_amount, prev_total, len(orders)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True)
    parser.add_argument("--master", help="Master file path")
    parser.add_argument("--store", help="Filter by store name (partial match)")
    args = parser.parse_args()

    if args.master:
        global MASTER_FILE
        MASTER_FILE = args.master

    target_date = datetime.strptime(args.date, "%Y-%m-%d")
    year, month, day = target_date.year, target_date.month, target_date.day

    print(f"\n--- Starting Full Invoice Regeneration for {args.date} ---")
    if args.store:
        print(f"[INFO] Store filter: {args.store}")

    store_totals, order_items, norm_to_original = load_order_data(MASTER_FILE, year, month)

    input_dir = os.path.join(PROJECT_ROOT, "data", "outputs", f"invoices_{year}{month:02d}{day:02d}")
    alt_input_dir = os.path.join(PROJECT_ROOT, "data", "outputs", "과거 파일", os.path.basename(input_dir))
    output_dir = os.path.join(PROJECT_ROOT, "data", "outputs", f"final_invoices_{year}{month:02d}{day:02d}")
    os.makedirs(output_dir, exist_ok=True)

    success_stores = []
    for norm_key in sorted(store_totals.keys()):
        if day not in store_totals[norm_key] or store_totals[norm_key][day] == 0:
            continue

        # 원본 매장명 복원 (드라이브 폴더 찾기, 템플릿 찾기용)
        store_name = norm_to_original.get(norm_key, norm_key)

        # --store 필터 적용
        if args.store and args.store not in store_name and args.store not in norm_key:
            continue
        print(f"Processing: {store_name} (key: {norm_key})")

        # Template - 여러 매장명 변형으로 검색
        template_path = None
        fn = f"{year}{month:02d}{day:02d}_{store_name.replace(' ', '_')}_거래명세서.xlsx"
        if os.path.exists(os.path.join(input_dir, fn)): template_path = os.path.join(input_dir, fn)
        elif os.path.exists(os.path.join(alt_input_dir, fn)): template_path = os.path.join(alt_input_dir, fn)
        else:
            # 정규화 키에서 공백 없는 버전으로도 검색
            fn_alt = f"{year}{month:02d}{day:02d}_{norm_key}_거래명세서.xlsx"
            if os.path.exists(os.path.join(input_dir, fn_alt)): template_path = os.path.join(input_dir, fn_alt)
            elif os.path.exists(os.path.join(alt_input_dir, fn_alt)): template_path = os.path.join(alt_input_dir, fn_alt)
            else: template_path = find_template(store_name, target_date)

        if not template_path:
            print(f"  [FAIL] No template for {store_name}")
            continue

        out_xlsx = os.path.join(output_dir, fn)
        res = process_statement(template_path, norm_key, day, out_xlsx, store_totals, order_items)

        if res:
            pdf_path = convert_to_pdf_win32(out_xlsx)

            # Drive
            dfolder = find_drive_folder(DRIVE_ROOT, store_name)
            dpath = os.path.join(DRIVE_ROOT, dfolder, f"{year}년 {month}월")
            os.makedirs(dpath, exist_ok=True)
            try:
                shutil.copy2(out_xlsx, os.path.join(dpath, os.path.basename(out_xlsx)))
                if pdf_path: shutil.copy2(pdf_path, os.path.join(dpath, os.path.basename(pdf_path)))
                print(f"  [OK] Uploaded to {dfolder}")
                success_stores.append(store_name)
            except Exception as e:
                print(f"  [ERROR] Drive upload failed: {e}")

    print(f"\nDone. Success: {len(success_stores)} stores.")

if __name__ == "__main__":
    main()
