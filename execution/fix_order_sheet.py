import openpyxl
import os
import datetime
import sys

# Add project root to path
project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(project_root)

# Import parsing logic
from execution.update_order_sheet import parse_processed_orders

path = r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터.xlsx"
sheet_name = "발주"

def fix_sheet():
    print(f"Loading {path}...")
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet_name]
    
    # 1. Identify rows to delete (Feb 02 rows)
    rows_to_delete = []
    max_r = ws.max_row
    print(f"Scanning {max_r} rows for cleanup...")
    
    # Scan backward for efficiency, but we need to target specific rows
    # Let's scan the last 2000 rows
    start_scan = max(1, max_r - 2000)
    
    for r in range(start_scan, max_r + 1):
        val_date = ws.cell(row=r, column=4).value
        # Check if date matches 2026-02-02
        matches = False
        if isinstance(val_date, datetime.datetime):
            if val_date.strftime("%Y-%m-%d") == "2026-02-02":
                matches = True
        elif isinstance(val_date, str):
            if "02월 02일" in val_date or "2026-02-02" in val_date:
                matches = True
                
        if matches:
            rows_to_delete.append(r)
            
    if rows_to_delete:
        print(f"Found {len(rows_to_delete)} rows to delete (2026-02-02).")
        # Delete from bottom up
        rows_to_delete.sort(reverse=True)
        # Use a more efficient deletion block strategy if contiguous?
        # For safety, just delete one by one or blocks.
        # openpyxl delete_rows is slow, so block deletion is better.
        
        # Group into ranges
        ranges = []
        if rows_to_delete:
            current_range = [rows_to_delete[0], rows_to_delete[0]]
            for r in rows_to_delete[1:]:
                if r == current_range[1] - 1:
                    current_range[1] = r
                else:
                    ranges.append(current_range)
                    current_range = [r, r]
            ranges.append(current_range)
            
        print(f"Deleting ranges: {ranges}")
        for end, start in ranges:
            count = end - start + 1
            print(f"Deleting {count} rows at {start}")
            ws.delete_rows(start, count)
    else:
        print("No rows found to delete.")

    # 2. Find insertion point (After Feb 01)
    # Re-calculate max row
    max_r = ws.max_row
    last_data_row = 0
    print("Finding insertion point...")
    for r in range(max_r, 1, -1):
        # Check existence of date or store name
        val_store = ws.cell(row=r, column=1).value
        val_date = ws.cell(row=r, column=4).value
        
        if val_store is not None or val_date is not None:
            last_data_row = r
            break
            
    start_row = last_data_row + 1
    print(f"Last data found at {last_data_row}. Insertion starts at {start_row}.")
    
    # 3. Parse and Insert
    orders = parse_processed_orders()
    if not orders:
        print("No orders to insert!")
        return

    print(f"Inserting {len(orders)} orders...")
    ws.insert_rows(start_row, amount=len(orders))
    
    # Styles
    from openpyxl.styles import Font, Border, Side, Alignment
    font_malgun = Font(name='Malgun Gothic', sz=11)
    font_calibri = Font(name='Calibri', sz=11)
    align_center = Alignment(horizontal='center', vertical='center')
    border_double_bottom = Border(bottom=Side(style='double'))
    border_none = Border()
    date_format = 'mm"월" dd"일"'
    price_format = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@'
    rate_format = '0.0'
    
    target_date_obj = datetime.datetime(2026, 2, 2)
    year = 2026
    month = 2
    
    current_row = start_row
    
    # Prepare sorting logic
    from collections import defaultdict
    grouped_orders = defaultdict(list)
    all_store_names = set()
    for order in orders:
        s_name = order['매장명']
        all_store_names.add(s_name)
        grouped_orders[s_name].append(order)
        
    shabu_stores = sorted([s for s in all_store_names if "샤브야키" in s])
    owl_stores = sorted([s for s in all_store_names if "부엉이산장" in s])
    dongwon_stores = sorted([s for s in all_store_names if "동원1차" in s])
    other_stores = sorted([s for s in all_store_names if "샤브야키" not in s and "부엉이산장" not in s and "동원1차" not in s])
    store_order = shabu_stores + owl_stores + other_stores + dongwon_stores
    
    for s_name in store_order:
        store_items = grouped_orders[s_name]
        store_start_row = current_row
        store_end_row = current_row + len(store_items) - 1
        
        profit_formula = (
            f"=IF(SUM(J{store_start_row}:J{store_end_row})>0,"
            f"(SUM(L{store_start_row}:L{store_end_row})-SUM(J{store_start_row}:J{store_end_row}))/SUM(J{store_start_row}:J{store_end_row})*100,0)"
        )
        
        for idx, order in enumerate(store_items):
            profit_val = profit_formula if idx == 0 else ""
            
            cells_data = {
                1: (order['매장명'], font_malgun, 'General'),
                2: (year, font_calibri, 'General'),
                3: (month, font_calibri, 'General'),
                4: (target_date_obj, font_calibri, date_format),
                5: (order['품목'], font_malgun, 'General'),
                7: (order['단위'], font_malgun, 'General'),
                8: (order['입고수량'], font_malgun, 'General'),
                9: (None, font_calibri, price_format),
                10: (f"=H{current_row}*I{current_row}", font_calibri, price_format),
                11: (None, font_calibri, price_format),
                12: (f"=K{current_row}*H{current_row}", font_calibri, price_format),
                13: (f"=IF(L{current_row}>0,(L{current_row}-J{current_row})/L{current_row}*100,0)", font_calibri, rate_format),
                14: (profit_val, font_calibri, rate_format)
            }
            
            for col, (val, font, fmt) in cells_data.items():
                cell = ws.cell(row=current_row, column=col, value=val)
                cell.font = font
                cell.alignment = align_center
                cell.number_format = fmt
                cell.border = border_none
                
            # Empty cols
            for col in [6, 15]:
                try:
                    cell = ws.cell(row=current_row, column=col, value=None)
                    cell.alignment = align_center
                    cell.border = border_none
                except: pass
                
            current_row += 1
            
        # Border
        for col in range(1, 15):
             ws.cell(row=store_end_row, column=col).border = border_double_bottom

    print("Saving workbook...")
    wb.save(path)
    print("Done.")

if __name__ == "__main__":
    fix_sheet()
