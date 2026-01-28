from docx import Document
import os

def analyze_docx_to_file(file_path, output_file):
    try:
        with open(output_file, 'a', encoding='utf-8') as f:
            f.write(f"\n\n=== Analyzing: {file_path} ===\n")
            
            if not os.path.exists(file_path):
                f.write("Error: File not found.\n")
                return

            doc = Document(file_path)
            
            f.write("--- Paragraphs ---\n")
            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    f.write(f"[P{i}] {para.text}\n")
            
            f.write("\n--- Tables ---\n")
            for i, table in enumerate(doc.tables):
                f.write(f"[Table {i}] R:{len(table.rows)} C:{len(table.columns)}\n")
                for r_idx, row in enumerate(table.rows):
                    row_cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    f.write(f"  R{r_idx}: {row_cells}\n")

    except Exception as e:
        with open(output_file, 'a', encoding='utf-8') as f:
            f.write(f"Error: {e}\n")

if __name__ == "__main__":
    out_file = "docx_struct_log.txt"
    if os.path.exists(out_file):
        os.remove(out_file)
        
    # 파일 1
    analyze_docx_to_file(r"c:\Users\DoKH_D\OneDrive\Desktop\일일 직납 현황\일일_직납_현황_260128.docx", out_file)
    # 파일 2
    analyze_docx_to_file(r"c:\Users\DoKH_D\OneDrive\Desktop\일일 직납 현황\일일_직납_현황_260127_1.docx", out_file)
