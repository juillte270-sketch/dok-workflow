import openpyxl
import os

path = r"G:\내 드라이브\1. 도크_주문 명세서\0. 매입단가_자료\도크발주관리데이터-AI접근용.xlsx"

try:
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb['발주']
    
    # We suspect data around 19300-19350 (Jan 27/28)
    for r in range(19300, 19350):
        # Check if this row has a store name in Col A (1)
        store = ws.cell(r, 1).value
        col_m = ws.cell(r, 13).value
        col_n = ws.cell(r, 14).value
        
        if store or col_m or col_n:
            print(f"Row {r}: Store={store}, M={col_m}, N={col_n}")
            
except Exception as e:
    print(e)
